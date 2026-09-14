# -*- coding: utf-8 -*-
"""
generate_presentation.py

Generates Class_06_Variables_and_Data_Types.pptx for
Bong Study Hub - Foundation Batch 2026.

Source of truth (do not redesign the lesson when editing this script):
  01_Master_Instructor_Guide/Class_06_Master_Instructor_Guide.md
  02_Student_Notes/Class_06_Student_Notes.md

This deck is a VISUAL TEACHING AID for a live instructor, not the
Student Notes placed onto slides. Slides carry diagrams, short
statements, questions, and live code the instructor points to and talks
around - never paragraphs read aloud.

Design system, component library, and file structure follow the
convention established in Classes 01-05's generate_presentation.py (same
palette, fonts, canvas, and shape-drawing primitives). This class reuses
Class 05's code-editor/output/error diagram toolkit directly wherever
the shape fits (print(name) vs. print("name") reuses the same
side-by-side code+output comparison Class 05 used for its two error
examples), and adds only the few new shapes Class 06 actually needs: a
single labeled box (the "variable as a labeled container" analogy) and
a two-panel labeled-code comparison (valid vs. invalid variable names).

Structure of this file:
  1. Theme configuration (colors, fonts, spacing)
  2. Low-level utility functions (shapes, text, arrows)
  3. Reusable slide-layout components (title, question, diagram, ...)
  4. Content-aware diagram builders (pipeline, vertical flow, code, ...)
  5. Slide-specific builders (slide_01 .. slide_22)
  6. Main generation function
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR

# ============================================================================
# 1. THEME CONFIGURATION
# ============================================================================

# -- Colors ---------------------------------------------------------------
# Same palette as Classes 01-05, for cross-class brand consistency.
NAVY = RGBColor(0x1E, 0x3A, 0x72)
BLUE = RGBColor(0x2A, 0x5C, 0xB8)
CYAN = RGBColor(0x22, 0xC7, 0xE0)
INK = RGBColor(0x1F, 0x27, 0x33)
MUTED = RGBColor(0x64, 0x70, 0x85)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_BG = RGBColor(0xF7, 0xF9, 0xFC)
CARD_BG = RGBColor(0xEE, 0xF3, 0xFB)
CARD_BORDER = RGBColor(0xC7, 0xD6, 0xEC)
CYAN_BG = RGBColor(0xE4, 0xF9, 0xFB)
CYAN_BORDER = RGBColor(0x9F, 0xE6, 0xEE)

# Carried over from Class 05: an error/warning accent, reused here only
# for labeling "invalid" variable names (Section 14 of the Master
# Guide) - not for a runtime error message this time.
ERROR_RED = RGBColor(0xB8, 0x3A, 0x3A)
ERROR_BG = RGBColor(0xFC, 0xEC, 0xEC)
ERROR_BORDER = RGBColor(0xEE, 0xB9, 0xB9)

# -- Fonts ------------------------------------------------------------------
FONT_DISPLAY = "Aptos Display"
FONT_BODY = "Aptos"
# Carried over from Class 05: a monospace face for anything that is
# literal Python source code or program output.
FONT_MONO = "Consolas"

# -- Canvas -----------------------------------------------------------------
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
MARGIN = Inches(0.7)
CONTENT_W = SLIDE_W - 2 * MARGIN

FOOTER_TEXT = "Bong Study Hub  ·  Foundation Batch 2026  ·  Class 06"


# ============================================================================
# 2. LOW-LEVEL UTILITIES
# ============================================================================

def new_slide(prs, bg=LIGHT_BG):
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = bg
    return slide


def _i(v):
    return int(round(v))


def _no_shadow(shape):
    try:
        shape.shadow.inherit = False
    except Exception:
        pass


def add_rect(slide, x, y, w, h, fill=WHITE, line=None, line_w=Pt(1),
             radius=0.06, shadow=False):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shp = slide.shapes.add_shape(shape_type, _i(x), _i(y), _i(w), _i(h))
    if radius:
        try:
            shp.adjustments[0] = radius
        except Exception:
            pass
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = line_w
    if not shadow:
        _no_shadow(shp)
    shp.text_frame.margin_left = 0
    shp.text_frame.margin_right = 0
    shp.text_frame.margin_top = 0
    shp.text_frame.margin_bottom = 0
    return shp


def add_oval(slide, x, y, w, h, fill=WHITE, line=None, line_w=Pt(1)):
    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, _i(x), _i(y), _i(w), _i(h))
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = line_w
    _no_shadow(shp)
    return shp


def add_text(slide, x, y, w, h, text, size=18, color=INK, bold=False,
             italic=False, align=PP_ALIGN.LEFT, font=FONT_BODY,
             anchor=MSO_ANCHOR.TOP, line_spacing=1.08, wrap=True,
             shrink=False):
    box = slide.shapes.add_textbox(_i(x), _i(y), _i(w), _i(h))
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    if shrink:
        tf.auto_size = MSO_AUTO_SIZE.NONE
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.name = font
        run.font.color.rgb = color
    return box


def add_line(slide, x1, y1, x2, y2, color=BLUE, width=Pt(1.5)):
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, _i(x1), _i(y1),
                                        _i(x2), _i(y2))
    conn.line.color.rgb = color
    conn.line.width = width
    _no_shadow(conn)
    return conn


def add_triangle(slide, cx, cy, size, color, rotation):
    s = _i(size)
    shp = slide.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE,
                                  _i(cx) - s // 2, _i(cy) - s // 2, s, s)
    shp.rotation = rotation
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.line.fill.background()
    _no_shadow(shp)
    return shp


def add_arrow_h(slide, x1, y, x2, color=BLUE, width=Pt(1.5)):
    tip_gap = Emu(90000)
    add_line(slide, x1, y, x2 - tip_gap, y, color=color, width=width)
    add_triangle(slide, x2 - Emu(45000), y, Emu(160000), color, rotation=90)


def add_arrow_v(slide, x, y1, y2, color=BLUE, width=Pt(1.5)):
    tip_gap = Emu(90000)
    add_line(slide, x, y1, x, y2 - tip_gap, color=color, width=width)
    add_triangle(slide, x, y2 - Emu(45000), Emu(160000), color, rotation=180)


def add_footer(slide):
    add_text(slide, MARGIN, SLIDE_H - Inches(0.42), Inches(7.5), Inches(0.3),
              FOOTER_TEXT, size=9.5, color=MUTED, font=FONT_BODY)


def add_page_number(slide, number):
    add_text(slide, SLIDE_W - Inches(1.1), SLIDE_H - Inches(0.42),
              Inches(0.5), Inches(0.3), str(number), size=9.5, color=MUTED,
              align=PP_ALIGN.RIGHT, font=FONT_BODY)


def add_kicker(slide, x, y, w, text, color=BLUE):
    add_text(slide, x, y, w, Inches(0.4), text.upper(), size=13, color=color,
              bold=True, font=FONT_BODY)


def set_notes(slide, purpose=None, say=None, question=None, expected=None,
               board=None, transition=None, timing=None, markers=None):
    lines = []
    if markers:
        lines.append("[" + "] [".join(markers) + "]")
    if timing:
        lines.append(f"TIME: {timing}")
    if purpose:
        lines.append(f"Purpose: {purpose}")
    if say:
        lines.append(f"Say: {say}")
    if question:
        lines.append(f"Ask: {question}")
    if expected:
        lines.append(f"Expected: {expected}")
    if board:
        lines.append(f"Board: {board}")
    if transition:
        lines.append(f"Transition: {transition}")
    slide.notes_slide.notes_text_frame.text = "\n".join(lines)


# ============================================================================
# 3. REUSABLE SLIDE-LAYOUT COMPONENTS
# ============================================================================

def layout_title(prs, kicker, title_lines, subtitle):
    slide = new_slide(prs, bg=NAVY)
    add_rect(slide, MARGIN, Inches(2.15), Inches(1.2), Inches(0.06), fill=CYAN)
    add_kicker(slide, MARGIN, Inches(1.75), Inches(10), kicker, color=CYAN)
    y = Inches(2.4)
    for line in title_lines:
        add_text(slide, MARGIN, y, Inches(11.5), Inches(1.1), line, size=50,
                  color=WHITE, bold=True, font=FONT_DISPLAY)
        y += Inches(1.0)
    add_text(slide, MARGIN, y + Inches(0.15), Inches(10.5), Inches(0.6),
              subtitle, size=20, color=RGBColor(0xB9, 0xCC, 0xEE),
              italic=True, font=FONT_BODY)
    add_text(slide, MARGIN, SLIDE_H - Inches(0.9), Inches(6), Inches(0.5),
              "Bong Study Hub\nFoundation Batch 2026", size=13,
              color=RGBColor(0x8F, 0xA6, 0xD1), font=FONT_BODY,
              line_spacing=1.15)
    return slide


def layout_section_divider(prs, kicker, title, subtitle=None, title_size=40):
    slide = new_slide(prs, bg=NAVY)
    add_rect(slide, MARGIN, Inches(3.05), Inches(1.0), Inches(0.06), fill=CYAN)
    if kicker:
        add_kicker(slide, MARGIN, Inches(2.65), Inches(10), kicker, color=CYAN)
    add_text(slide, MARGIN, Inches(3.25), Inches(11.8), Inches(2.4), title,
              size=title_size, color=WHITE, bold=True, font=FONT_DISPLAY,
              line_spacing=1.15)
    if subtitle:
        add_text(slide, MARGIN, Inches(5.6), Inches(10.5), Inches(0.8),
                  subtitle, size=18, color=RGBColor(0xB9, 0xCC, 0xEE),
                  font=FONT_BODY, italic=True)
    add_text(slide, MARGIN, SLIDE_H - Inches(0.42), Inches(7.5), Inches(0.3),
              FOOTER_TEXT, size=9.5, color=RGBColor(0x6B, 0x82, 0xAE))
    return slide


def _content_header(slide, kicker, title, title_size=32):
    if kicker:
        add_kicker(slide, MARGIN, Inches(0.5), CONTENT_W, kicker, color=BLUE)
        title_y = Inches(0.86)
    else:
        title_y = Inches(0.6)
    add_text(slide, MARGIN, title_y, CONTENT_W, Inches(0.9), title,
              size=title_size, color=NAVY, bold=True, font=FONT_DISPLAY)
    add_rect(slide, MARGIN, title_y + Inches(0.85), Inches(0.85),
              Inches(0.045), fill=CYAN)
    return title_y + Inches(1.15)


def layout_big_question(prs, kicker, question, note=None, size=38):
    slide = new_slide(prs, bg=LIGHT_BG)
    if kicker:
        add_kicker(slide, MARGIN, Inches(1.5), CONTENT_W, kicker, color=BLUE)
    add_text(slide, Inches(1.0), Inches(2.3), Inches(11.3), Inches(2.9),
              question, size=size, color=NAVY, bold=True, font=FONT_DISPLAY,
              align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
              line_spacing=1.15)
    if note:
        add_text(slide, Inches(1.0), Inches(5.5), Inches(11.3), Inches(0.5),
                  note, size=15, color=MUTED, align=PP_ALIGN.CENTER,
                  italic=True)
    add_footer(slide)
    return slide


def layout_two_column(prs, kicker, title, left, right, connector_label=None,
                       connector_symbol=None):
    slide = new_slide(prs, bg=LIGHT_BG)
    y = _content_header(slide, kicker, title)
    col_w = (CONTENT_W - Inches(0.6)) / 2
    card_h = Inches(2.7)
    card_y = y + Inches(0.4)
    for i, card in enumerate((left, right)):
        x = MARGIN + i * (col_w + Inches(0.6))
        fill = CARD_BG if i == 0 else CYAN_BG
        border = CARD_BORDER if i == 0 else CYAN_BORDER
        add_rect(slide, x, card_y, col_w, card_h, fill=fill, line=border,
                  line_w=Pt(1.25), radius=0.05)
        if card.get("tag"):
            add_text(slide, x + Inches(0.35), card_y + Inches(0.3),
                      col_w - Inches(0.7), Inches(0.4), card["tag"],
                      size=13, color=(BLUE if i == 0 else NAVY), bold=True)
        add_text(slide, x + Inches(0.35), card_y + Inches(0.75),
                  col_w - Inches(0.7), Inches(0.7), card["title"],
                  size=card.get("title_size", 25), color=NAVY, bold=True,
                  font=card.get("title_font", FONT_DISPLAY))
        add_text(slide, x + Inches(0.35), card_y + Inches(1.55),
                  col_w - Inches(0.7), Inches(1.0), card["sub"], size=16,
                  color=INK, line_spacing=1.2)
    mid_y = card_y + card_h / 2
    gap_x1 = MARGIN + col_w
    gap_x2 = gap_x1 + Inches(0.6)
    if connector_symbol:
        add_text(slide, gap_x1, mid_y - Inches(0.35), gap_x2 - gap_x1,
                  Inches(0.7), connector_symbol, size=30, color=NAVY,
                  bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
                  font=FONT_DISPLAY)
    else:
        add_arrow_h(slide, gap_x1 + Inches(0.08), mid_y, gap_x2 - Inches(0.08),
                    color=BLUE, width=Pt(2))
        if connector_label:
            add_text(slide, gap_x1, mid_y - Inches(0.55), Inches(0.6),
                      Inches(0.35), connector_label, size=11, color=MUTED,
                      align=PP_ALIGN.CENTER)
    add_footer(slide)
    return slide


def layout_diagram(prs, kicker, title, diagram_fn, caption=None,
                    title_size=32):
    slide = new_slide(prs, bg=LIGHT_BG)
    y = _content_header(slide, kicker, title, title_size=title_size)
    area_y = y + Inches(0.25)
    area_h = SLIDE_H - area_y - Inches(0.75)
    if caption:
        area_h -= Inches(0.5)
    diagram_fn(slide, MARGIN, area_y, CONTENT_W, area_h)
    if caption:
        add_text(slide, MARGIN, SLIDE_H - Inches(1.05), CONTENT_W,
                  Inches(0.4), caption, size=13, color=MUTED, italic=True,
                  align=PP_ALIGN.CENTER)
    add_footer(slide)
    return slide


def layout_full_screen_visual(prs, kicker, title, diagram_fn, caption=None):
    slide = new_slide(prs, bg=LIGHT_BG)
    add_kicker(slide, MARGIN, Inches(0.45), CONTENT_W, kicker, color=BLUE)
    if title:
        add_text(slide, MARGIN, Inches(0.8), CONTENT_W, Inches(0.7), title,
                  size=26, color=NAVY, bold=True, font=FONT_DISPLAY)
    diagram_fn(slide, MARGIN, Inches(1.55), CONTENT_W,
                SLIDE_H - Inches(1.55) - Inches(1.15))
    if caption:
        add_text(slide, MARGIN, SLIDE_H - Inches(1.0), CONTENT_W, Inches(0.4),
                  caption, size=13, color=MUTED, italic=True,
                  align=PP_ALIGN.CENTER)
    add_footer(slide)
    return slide


def layout_activity(prs, kicker, title, prompt_lines, sub=None, tag="ACTIVITY"):
    slide = new_slide(prs, bg=LIGHT_BG)
    add_rect(slide, Inches(0.9), Inches(0.55), Inches(1.5), Inches(0.42),
              fill=CYAN_BG, line=CYAN_BORDER, line_w=Pt(1))
    add_text(slide, Inches(0.9), Inches(0.55), Inches(1.5), Inches(0.42),
              tag, size=13, color=NAVY, bold=True, align=PP_ALIGN.CENTER,
              anchor=MSO_ANCHOR.MIDDLE)
    add_kicker(slide, Inches(2.6), Inches(0.62), Inches(9), kicker,
                color=BLUE)
    add_text(slide, Inches(0.9), Inches(1.25), Inches(11.5), Inches(1.0),
              title, size=30, color=NAVY, bold=True, font=FONT_DISPLAY)
    y = Inches(2.5)
    for line in prompt_lines:
        add_text(slide, Inches(0.9), y, Inches(11.5), Inches(0.85), line,
                  size=24, color=INK, font=FONT_BODY, line_spacing=1.15)
        y += Inches(0.8)
    if sub:
        add_text(slide, Inches(0.9), y + Inches(0.15), Inches(11.0),
                  Inches(0.6), sub, size=15, color=MUTED, italic=True)
    add_footer(slide)
    return slide


# ============================================================================
# 4. CONTENT-AWARE DIAGRAM BUILDERS
# ============================================================================

def draw_pipeline(items, height=Inches(1.3), font_size=15,
                   highlight_indices=None, gap=Inches(0.55)):
    highlight_indices = highlight_indices or set()

    def _fn(slide, x, y, w, h):
        n = len(items)
        box_w = (w - gap * (n - 1)) / n
        box_y = y + (h - height) / 2
        cx = x
        for i, (label, sub) in enumerate(items):
            is_hl = i in highlight_indices
            fill = CYAN_BG if is_hl else CARD_BG
            border = CYAN_BORDER if is_hl else CARD_BORDER
            add_rect(slide, cx, box_y, box_w, height, fill=fill,
                      line=border, line_w=Pt(1.25), radius=0.09)
            label_y = box_y + (height / 2 - Inches(0.32) if sub
                                else height / 2 - Inches(0.24))
            add_text(slide, cx, label_y, box_w, Inches(0.5), label,
                      size=font_size, color=NAVY, bold=True,
                      align=PP_ALIGN.CENTER, font=FONT_BODY)
            if sub:
                add_text(slide, cx, label_y + Inches(0.42), box_w,
                          Inches(0.35), sub, size=font_size - 4, color=MUTED,
                          align=PP_ALIGN.CENTER)
            if i < n - 1:
                add_arrow_h(slide, cx + box_w + Inches(0.06), box_y + height / 2,
                            cx + box_w + gap - Inches(0.06), color=BLUE,
                            width=Pt(1.75))
            cx += box_w + gap
    return _fn


def draw_vertical_flow(items, highlight_indices=None, box_w=Inches(4.6),
                        font_size=16):
    highlight_indices = highlight_indices or set()

    def _fn(slide, x, y, w, h):
        n = len(items)
        box_h = min(Inches(0.75), h / (n * 1.35))
        gap = (h - n * box_h) / (n - 1) if n > 1 else Inches(0)
        gap = max(Inches(0.1), min(gap, Inches(0.5)))
        total_h = n * box_h + (n - 1) * gap
        cy = y + max(Inches(0), (h - total_h) / 2)
        bx = x + w / 2 - box_w / 2
        for i, label in enumerate(items):
            is_hl = i in highlight_indices
            fill = CYAN_BG if is_hl else CARD_BG
            border = CYAN_BORDER if is_hl else CARD_BORDER
            add_rect(slide, bx, cy, box_w, box_h, fill=fill, line=border,
                      line_w=Pt(1.4 if is_hl else 1.1), radius=0.14)
            add_text(slide, bx, cy, box_w, box_h, label, size=font_size,
                      color=NAVY, bold=True, align=PP_ALIGN.CENTER,
                      anchor=MSO_ANCHOR.MIDDLE)
            if i < n - 1:
                add_arrow_v(slide, x + w / 2, cy + box_h + Inches(0.03),
                            cy + box_h + gap - Inches(0.03), color=BLUE,
                            width=Pt(1.75))
            cy += box_h + gap
    return _fn


def draw_word_grid(words, cols=4, highlight_indices=None, font_size=20):
    highlight_indices = highlight_indices or set()

    def _fn(slide, x, y, w, h):
        n = len(words)
        rows = (n + cols - 1) // cols
        gap = Inches(0.3)
        cell_w = (w - gap * (cols - 1)) / cols
        cell_h = min(Inches(1.0), (h - gap * (rows - 1)) / rows) if rows else h
        grid_h = rows * cell_h + (rows - 1) * gap
        top = y + max(Inches(0), (h - grid_h) / 2)
        for i, word in enumerate(words):
            r, c = divmod(i, cols)
            last_row = (r == rows - 1)
            items_in_row = cols if (not last_row or n % cols == 0) else n % cols
            row_w = items_in_row * cell_w + (items_in_row - 1) * gap
            row_x0 = x + (w - row_w) / 2
            cx = row_x0 + c * (cell_w + gap)
            cy = top + r * (cell_h + gap)
            is_hl = i in highlight_indices
            fill = CYAN_BG if is_hl else CARD_BG
            border = CYAN_BORDER if is_hl else CARD_BORDER
            add_rect(slide, cx, cy, cell_w, cell_h, fill=fill, line=border,
                      line_w=Pt(1.25), radius=0.14)
            add_text(slide, cx, cy, cell_w, cell_h, word, size=font_size,
                      color=NAVY, bold=True, align=PP_ALIGN.CENTER,
                      anchor=MSO_ANCHOR.MIDDLE, font=FONT_BODY)
    return _fn


def draw_definition(term, definition, term_size=72, def_size=28):
    def _fn(slide, x, y, w, h):
        add_text(slide, x, y, w, h * 0.42, term, size=term_size, color=NAVY,
                  bold=True, font=FONT_DISPLAY, align=PP_ALIGN.CENTER,
                  anchor=MSO_ANCHOR.MIDDLE)
        rule_y = y + h * 0.46
        add_rect(slide, x + w / 2 - Inches(0.5), rule_y, Inches(1.0),
                  Inches(0.04), fill=CYAN)
        add_text(slide, x, rule_y + Inches(0.25), w, h * 0.5 - Inches(0.25),
                  definition, size=def_size, color=INK, bold=True,
                  align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.TOP,
                  line_spacing=1.2, font=FONT_BODY)
    return _fn


def draw_statement(text, size=40, color=NAVY):
    def _fn(slide, x, y, w, h):
        add_text(slide, x, y, w, h, text, size=size, color=color, bold=True,
                  font=FONT_DISPLAY, align=PP_ALIGN.CENTER,
                  anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.15)
    return _fn


def draw_table(headers, rows, highlight_cols=None, font_size=18):
    highlight_cols = highlight_cols or set()

    def _fn(slide, x, y, w, h):
        ncols = len(headers)
        nrows = len(rows)
        row_h = min(Inches(0.85), h / (nrows + 1))
        total_h = row_h * (nrows + 1)
        top = y + max(Inches(0), (h - total_h) / 2)
        col_w = w / ncols
        for c, htext in enumerate(headers):
            cx = x + c * col_w
            add_rect(slide, cx + Inches(0.03), top + Inches(0.03),
                      col_w - Inches(0.06), row_h - Inches(0.06), fill=NAVY,
                      radius=0.08)
            add_text(slide, cx, top, col_w, row_h, htext, size=font_size,
                      color=WHITE, bold=True, align=PP_ALIGN.CENTER,
                      anchor=MSO_ANCHOR.MIDDLE)
        for r, row in enumerate(rows):
            ry = top + (r + 1) * row_h
            for c, val in enumerate(row):
                cx = x + c * col_w
                is_hl = c in highlight_cols
                fill = CYAN_BG if is_hl else (WHITE if r % 2 == 0 else CARD_BG)
                border = CYAN_BORDER if is_hl else CARD_BORDER
                add_rect(slide, cx + Inches(0.03), ry + Inches(0.03),
                          col_w - Inches(0.06), row_h - Inches(0.06),
                          fill=fill, line=border, line_w=Pt(1), radius=0.06)
                add_text(slide, cx, ry, col_w, row_h, str(val),
                          size=font_size, color=INK, align=PP_ALIGN.CENTER,
                          anchor=MSO_ANCHOR.MIDDLE,
                          font=(FONT_MONO if c == 1 else FONT_BODY))
    return _fn


# -- Carried over from Class 05: code / output shapes -----------------------

def draw_code_block(lines, font_size=22):
    """A dark 'editor' card holding monospace Python source, one string per
    line. The small cyan dot in the corner reads as a window/tab marker."""
    def _fn(slide, x, y, w, h):
        n = max(1, len(lines))
        card_h = min(h, Inches(0.62) * n + Inches(0.55))
        top = y + max(Inches(0), (h - card_h) / 2)
        add_rect(slide, x, top, w, card_h, fill=NAVY, radius=0.06)
        add_oval(slide, x + Inches(0.32), top + Inches(0.28), Inches(0.14),
                  Inches(0.14), fill=CYAN, line=None)
        line_h = (card_h - Inches(0.55)) / n
        ty = top + Inches(0.48)
        for i, line in enumerate(lines):
            add_text(slide, x + Inches(0.62), ty + i * line_h,
                      w - Inches(0.95), line_h, line, size=font_size,
                      color=WHITE, font=FONT_MONO, anchor=MSO_ANCHOR.MIDDLE)
    return _fn


def draw_output_card(lines, is_error=False, font_size=20, label=None):
    """A light 'console' card for program output, or an error-tinted card
    when is_error is True."""
    def _fn(slide, x, y, w, h):
        fill = ERROR_BG if is_error else CARD_BG
        border = ERROR_BORDER if is_error else CARD_BORDER
        bar = ERROR_RED if is_error else CYAN
        text_color = ERROR_RED if is_error else INK
        add_rect(slide, x, y, w, h, fill=fill, line=border, line_w=Pt(1.25),
                  radius=0.05)
        add_rect(slide, x, y, Inches(0.09), h, fill=bar, radius=None)
        lab = label or ("ERROR" if is_error else "OUTPUT")
        add_text(slide, x + Inches(0.35), y + Inches(0.16),
                  w - Inches(0.6), Inches(0.32), lab, size=12.5,
                  color=(ERROR_RED if is_error else MUTED), bold=True)
        n = max(1, len(lines))
        body_y = y + Inches(0.55)
        body_h = h - Inches(0.7)
        add_text(slide, x + Inches(0.35), body_y, w - Inches(0.6), body_h,
                  "\n".join(lines), size=font_size, color=text_color,
                  font=FONT_MONO, anchor=MSO_ANCHOR.TOP, line_spacing=1.3)
    return _fn


def draw_code_and_output(code_lines, output_lines, is_error=False,
                          code_font=22, output_font=18, output_label=None):
    """Editor card on top, console result (or error) beneath it."""
    def _fn(slide, x, y, w, h):
        gap = Inches(0.35)
        code_h = h * 0.54
        out_h = h - code_h - gap
        draw_code_block(code_lines, font_size=code_font)(slide, x, y, w, code_h)
        draw_output_card(output_lines, is_error=is_error,
                          font_size=output_font, label=output_label)(
            slide, x, y + code_h + gap, w, out_h)
    return _fn


def draw_two_code_examples(left, right, font_size=17, output_font=15):
    """Two side-by-side code+output panels - reused here from Class 05 for
    the print(name) vs. print("name") contrast (Master Guide Section 12),
    the single most important visual moment of this class."""
    def _fn(slide, x, y, w, h):
        gap = Inches(0.55)
        col_w = (w - gap) / 2
        for i, spec in enumerate((left, right)):
            cx = x + i * (col_w + gap)
            draw_code_and_output(
                spec["code"], spec["output"], is_error=spec.get("is_error", False),
                code_font=font_size, output_font=output_font,
                output_label=spec.get("label"))(slide, cx, y, col_w, h)
    return _fn


# -- New for Class 06: labeled-box and labeled-code shapes ------------------

def draw_labeled_box(label, value, label_size=18, value_size=46):
    """The 'variable as a labeled container' analogy (Master Guide
    Section 10): a small tab bearing the variable's name, sitting on a
    box holding its current value."""
    def _fn(slide, x, y, w, h):
        box_w = min(w, Inches(4.4))
        box_h = min(h - Inches(0.6), Inches(2.4))
        bx = x + (w - box_w) / 2
        by = y + (h - box_h) / 2 + Inches(0.3)
        tab_w = Inches(2.3)
        tab_h = Inches(0.55)
        tx = bx + (box_w - tab_w) / 2
        add_rect(slide, tx, by - tab_h + Inches(0.1), tab_w, tab_h,
                  fill=NAVY, radius=0.16)
        add_text(slide, tx, by - tab_h + Inches(0.1), tab_w, tab_h, label,
                  size=label_size, color=WHITE, bold=True,
                  align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
                  font=FONT_MONO)
        add_rect(slide, bx, by, box_w, box_h, fill=CYAN_BG, line=CYAN_BORDER,
                  line_w=Pt(1.75), radius=0.08)
        add_text(slide, bx, by, box_w, box_h, str(value), size=value_size,
                  color=NAVY, bold=True, align=PP_ALIGN.CENTER,
                  anchor=MSO_ANCHOR.MIDDLE, font=FONT_MONO)
    return _fn


def draw_two_labeled_code(left_label, left_lines, right_label, right_lines,
                           right_is_flagged=True, font_size=18):
    """Two labeled code cards side by side - used for valid vs. invalid
    variable names (Master Guide Section 14). The right label turns
    error-red when right_is_flagged, without implying a runtime crash -
    just 'this one breaks the naming rules.'"""
    def _fn(slide, x, y, w, h):
        gap = Inches(0.5)
        col_w = (w - gap) / 2
        label_h = Inches(0.5)
        code_h = h - label_h - Inches(0.15)
        specs = [(left_label, left_lines, False),
                 (right_label, right_lines, right_is_flagged)]
        for i, (label, lines, flagged) in enumerate(specs):
            cx = x + i * (col_w + gap)
            color = ERROR_RED if flagged else BLUE
            add_text(slide, cx, y, col_w, label_h, label, size=16,
                      color=color, bold=True, align=PP_ALIGN.CENTER)
            draw_code_block(lines, font_size=font_size)(
                slide, cx, y + label_h + Inches(0.1), col_w, code_h)
    return _fn


# ============================================================================
# 5. SLIDE-SPECIFIC BUILDERS
# ============================================================================

def slide_01_title(prs):
    slide = layout_title(
        prs, "Bong Study Hub · Foundation Batch 2026 · Class 06",
        ["Variables &", "Data Types"],
        "Giving Values a Label That Can Change")
    set_notes(slide, timing="2 min",
              purpose="Set tone; do not read the slide aloud.",
              say="Open with the live scoreboard demonstration (Guide "
                  "§8) before this slide is even shown, if possible.",
              transition="If a program needs to remember something that "
                          "changes, how does it hold onto it?")
    return slide


def slide_02_big_question(prs):
    slide = layout_big_question(
        prs, "The Big Question",
        "If a program needs to remember\nsomething that changes — a "
        "score, an age,\nan answer — how does it hold onto it?",
        size=30)
    set_notes(slide, markers=["ASK", "PAUSE", "DO NOT ANSWER YET"],
              timing="(closes Guide §8)",
              purpose="Land the central question right after the live "
                      "'edit and re-run the scoreboard' demonstration.",
              board="None.",
              transition="Let's go back to exactly where Class 05 left "
                          "off.")
    return slide


def slide_03_callback(prs):
    slide = layout_two_column(
        prs, "Two Classes, One Missing Piece", "Reopening Two Boxes",
        {"tag": "CLASS 05 LEFT OPEN", "title": "Fixed Text",
         "sub": "Every program said the exact same thing, every single "
                "time — changing it meant editing the code and running "
                "it again."},
        {"tag": "CLASS 04 REOPENED", "title": "Bits",
         "sub": "Information is represented inside a computer as bits — "
                "you never see them directly."},
        connector_symbol="+")
    set_notes(slide, markers=["ASK"], timing="5 min (Guide §9)",
              purpose="Re-open Class 05's unresolved hook and Class 04's "
                      "bits idea together - both get answered today.",
              question="What did Class 05 leave unanswered? What did "
                        "Class 04 teach about how a computer holds "
                        "information?",
              expected="How to hold onto something that changes; it "
                        "represents information as bits.",
              board="None.",
              transition="Let's name the fix.")
    return slide


def slide_04_what_is_a_variable(prs):
    slide = layout_diagram(
        prs, "Naming the Fix", "What Is a Variable?",
        draw_labeled_box("SCORE", "0"),
        caption='score = 0   →   print(score)   →   shows 0')
    set_notes(slide, timing="7 min (Guide §10)",
              purpose="Introduce the labeled-box analogy before any "
                      "code, then type score = 0 / print(score) live.",
              say="A variable is a named place to store a value - and "
                  "unlike Class 05's fixed text, its value can change "
                  "later. score is the label. 0 is what's currently "
                  "inside.",
              question="In the labeled-box idea, what's the label? "
                        "What's inside?",
              expected="The variable name is the label; the value is "
                        "what's inside.",
              board="Live screen.",
              transition="Let's look closely at the line that actually "
                          "created it.")
    return slide


def slide_05_assignment_statement(prs):
    slide = layout_diagram(
        prs, "Read It Correctly", "= Means Store, Not Equals",
        draw_statement('score = 0\n\nREAD AS: "score GETS 0"\nNOT: '
                        '"score equals 0"', size=30))
    set_notes(slide, timing="(opens Guide §11) — NON-NEGOTIABLE",
              purpose="Land the core reframe before proving it live - in "
                      "Python, = means store, not mathematical equality.",
              question="In math class, what does = mean?",
              expected="Equals — both sides are the same.",
              board="None.",
              transition="Let's prove this by actually changing what's "
                          "stored.")
    return slide


def slide_06_assignment_proof(prs):
    slide = layout_diagram(
        prs, "Proof, Live", "One Name, Two Different Values",
        draw_code_and_output(
            ["score = 0", "print(score)", "score = 10", "print(score)"],
            ["0", "10"], code_font=22, output_font=22))
    set_notes(slide, markers=["LIVE"], timing="(closes Guide §11)",
              purpose="If = meant mathematical equality, this would be a "
                      "contradiction. It isn't - each = just overwrites "
                      "what was there.",
              question="How should you read age = 18 out loud?",
              expected='"age gets 18," not "age equals 18."',
              board="Live screen.",
              transition="Now let's use this variable inside print() - "
                          "and see something that trips up almost every "
                          "beginner at least once.")
    return slide


def slide_07_print_name_vs_string(prs):
    slide = layout_diagram(
        prs, "The Most Important Contrast Today",
        "print(name)  vs.  print(\"name\")",
        draw_two_code_examples(
            {"code": ['name = "Priya"', "print(name)"],
             "output": ["Priya"], "label": "NO QUOTES → THE STORED VALUE"},
            {"code": ['name = "Priya"', 'print("name")'],
             "output": ["name"], "label": "QUOTES → THE LITERAL WORD"},
        ))
    set_notes(slide, markers=["TYPE-ALONG", "ASK"], timing="7 min "
              "(Guide §12) — NON-NEGOTIABLE",
              purpose="Every student types both lines and predicts "
                      "before running - the single most important "
                      "moment of the class.",
              say="Quotes are the entire difference. This is the "
                  "easiest mistake to make in the next several classes "
                  "- let's make it on purpose, right now.",
              question="Before I run these - will they show the same "
                        "thing?",
              expected="No - print(name) shows Priya; print(\"name\") "
                        "shows the literal word name.",
              board="Live screen.",
              transition="Let's change what's stored in a variable and "
                          "watch what happens.")
    return slide


def slide_08_reassignment(prs):
    slide = layout_diagram(
        prs, "Building Live", "Reassignment",
        draw_code_and_output(
            ["score = 0", "print(score)", "score = 10", "print(score)",
             "score = 25", "print(score)"],
            ["0", "10", "25"], code_font=17, output_font=19,
            output_label="OUTPUT — the old value is never seen again"))
    set_notes(slide, markers=["LIVE", "ASK"], timing="7 min (Guide §13) "
              "— NON-NEGOTIABLE",
              purpose="Build this live, one reassignment at a time - "
                      "the old value is completely overwritten, not "
                      "remembered.",
              question="After score = 10 runs, if I ask Python for "
                        "score again, what happened to the 0?",
              expected="It's completely gone - overwritten, not "
                        "remembered anywhere.",
              board="Live screen.",
              transition="Before we go further, let's talk about what "
                          "you're actually allowed to name a variable.")
    return slide


def slide_09_naming_variables(prs):
    slide = layout_diagram(
        prs, "The Rules", "Naming Variables",
        draw_two_labeled_code(
            "VALID", ['age = 18', 'student_name = "Rio"', "_temp = 5"],
            "INVALID", ["2nd_place = ...", "student name = ...",
                        "class = ..."],
            font_size=16),
        caption="Starts with a digit · contains a space · a reserved "
                "Python word — three different problems, all invalid.")
    set_notes(slide, timing="5 min (Guide §14)",
              purpose="State the hard rules, then show valid/invalid "
                      "side by side. Case-sensitivity matters too: age "
                      "and Age are different variables.",
              say="Beyond what's required, good style uses snake_case "
                  "and descriptive names - student_age beats x.",
              question="Can a variable name start with a digit? Are age "
                        "and Age the same variable?",
              expected="No. No - Python is case-sensitive.",
              board="Live screen.",
              transition="So far every variable held text or a number. "
                          "Let's name the different kinds of values a "
                          "variable can hold.")
    return slide


def slide_10_meet_data_types(prs):
    slide = layout_diagram(
        prs, "Four Kinds of Value", "Meet the Data Types",
        draw_table(
            ["Variable", "Value", "Type"],
            [["age", "18", "int"], ["price", "19.99", "float"],
             ["name", '"Priya"', "str"], ["is_student", "True", "bool"]],
            font_size=19),
        caption='print(type(age))  →  <class \'int\'>')
    set_notes(slide, markers=["LIVE"], timing="8 min (Guide §15) — "
              "NON-NEGOTIABLE",
              purpose="Type all four lines live, then name each type. "
                      "type() is a peek, not a deep dive.",
              question="What are the four data types we just met?",
              expected="int, float, str, bool.",
              board="Live screen.",
              transition="Let's slow down on the first one - numbers "
                          "actually come in two flavors.")
    return slide


def slide_11_numbers(prs):
    slide = layout_diagram(
        prs, "Two Flavors", "Numbers: int vs. float",
        draw_two_labeled_code(
            "INT — A WHOLE NUMBER", ["apples = 3"],
            "FLOAT — HAS A DECIMAL POINT", ["temperature = 36.6"],
            right_is_flagged=False, font_size=20))
    set_notes(slide, timing="6 min (Guide §16)",
              purpose="The decimal point is the entire test - not a "
                      "value judgment, just a distinction.",
              say="We're not adding, subtracting, or doing any math "
                  "with these numbers yet - that's next class's entire "
                  "job.",
              question="Is 7 an int or a float? Is 7.0 an int or a "
                        "float?",
              expected="int. Float - it has a decimal point.",
              board="Live screen.",
              transition="Strings we already know from Class 05 - but "
                          "let's revisit exactly where quotes do and "
                          "don't belong.")
    return slide


def slide_12_strings_revisited(prs):
    slide = layout_diagram(
        prs, "The Rule Hasn't Changed", "Strings, Revisited",
        draw_two_code_examples(
            {"code": ['city = "Kolkata"', "print(city)"],
             "output": ["Kolkata"], "label": "A VARIABLE — NO QUOTES"},
            {"code": ['city = "Kolkata"', 'print("city")'],
             "output": ["city"], "label": "LITERAL TEXT — QUOTES"},
        ))
    set_notes(slide, timing="5 min (Guide §17)",
              purpose="Confirm Section 12's lesson one more time, "
                      "explicitly tied to the Class 05 quotes rule.",
              question="Does a variable name ever go in quotes?",
              expected="No.",
              board="Live screen.",
              transition="One more data type - one that isn't a number "
                          "or text at all.")
    return slide


def slide_13_booleans(prs):
    slide = layout_diagram(
        prs, "Yes or No", "Booleans",
        draw_code_and_output(
            ["is_raining = True", "is_weekend = False",
             "print(is_raining)"],
            ["True"], code_font=20, output_font=20))
    set_notes(slide, timing="6 min (Guide §18)",
              purpose="True/False - capitalized, no quotes. Do not "
                      "introduce comparison operators here.",
              say='True and "True" are not the same thing. Without '
                  "quotes, it's a boolean. With quotes, it's just a "
                  "four-letter string.",
              question="What are the only two possible boolean values? "
                        "Do booleans need quotes?",
              expected="True and False. No.",
              board="Live screen.",
              transition="Let's put a variable inside a full sentence, "
                          "not just print it alone.")
    return slide


def slide_14_combining(prs):
    slide = layout_diagram(
        prs, "Building a Sentence", "Combining Text and Variables",
        draw_code_and_output(
            ['name = "Priya"', "age = 18",
             'print("My name is", name, "and I am", age, "years old.")'],
            ["My name is Priya and I am 18 years old."],
            code_font=15, output_font=17))
    set_notes(slide, markers=["LIVE", "ASK"], timing="7 min (Guide §19) "
              "— NON-NEGOTIABLE",
              purpose="A comma inside print() mixes literal text and "
                      "variables - Python adds a space automatically.",
              question="Before I run this - what do you think the "
                        "output will look like?",
              expected="A single readable sentence combining the text "
                        "and the two variables' values.",
              board="Live screen.",
              transition="Let's reconnect this whole idea to something "
                          "from several classes ago.")
    return slide


def slide_15_class04_callback(prs):
    slide = layout_diagram(
        prs, "Reconnecting", "A Variable Is a Labeled Representation",
        draw_vertical_flow(
            ["INFORMATION", "REPRESENTATION  (Class 04)", "BITS  (Class 04)",
             "VARIABLE  (today)"], highlight_indices={3}, box_w=Inches(5.6),
            font_size=17))
    set_notes(slide, timing="7 min (Guide §20)",
              purpose="A variable is the human-readable label Python "
                      "gives you for a piece of stored representation - "
                      "you never have to think about the bits directly.",
              question="What did Class 04 say information gets "
                        "represented as? What does a variable give you "
                        "that raw bits don't?",
              expected="Bits. A readable, meaningful name for the "
                        "stored value.",
              board="None.",
              transition="Let's put the entire day's reasoning on one "
                          "board - starting all the way back at Class "
                          "01.")
    return slide


def slide_16_full_bridge(prs):
    items = ["PROBLEM", "LOGIC", "ALGORITHM", "PYTHON CODE", "VARIABLES",
             "RUN", "OUTPUT / RESULT"]
    slide = layout_full_screen_visual(
        prs, "The Hero Diagram", None,
        draw_vertical_flow(items, highlight_indices={4}, box_w=Inches(5.2),
                            font_size=17),
        caption="Cyan = new today  ·  Card = already known from Classes "
                "01 and 05")
    set_notes(slide, markers=["LIVE"], timing="7 min (Guide §21) — "
              "NON-NEGOTIABLE, HERO MOMENT",
              purpose="Build this top to bottom, one line at a time - "
                      "never reveal it finished. Today inserted one new "
                      "box - VARIABLES - into Class 05's chain.",
              board="Whiteboard - the one non-screen drawing before the "
                    "activity.",
              transition="Time to try this yourselves.")
    return slide


def slide_17_activity_intro(prs):
    slide = layout_activity(
        prs, "Try It Yourself", "Variable Profile Card",
        ["At least four variables:", "one string, one int, one float, "
         "one boolean."],
        sub="Then print() each one in a full sentence, combining text "
            "and variables with commas.")
    set_notes(slide, markers=["ACTIVITY"], timing="14 min (Guide §22) "
              "— NON-NEGOTIABLE, HERO ACTIVITY",
              purpose="Every student extends Class 05's 'About Me' idea "
                      "using real variables and data types. Circulate "
                      "constantly.",
              say="Stuck? What's one true number about you? One true "
                  "yes/no fact?",
              board="Circulate and read screens rather than calling on "
                    "students verbally.",
              transition="Here's a template if you want a starting "
                          "point.")
    return slide


def slide_18_activity_template(prs):
    slide = layout_diagram(
        prs, "A Starting Point (Optional)", "Variable Profile Card — Template",
        draw_code_block(
            ['name = "___"', "age = ___", "height = ___",
             "is_student = ___", "", 'print("My name is", name)'],
            font_size=18))
    set_notes(slide, timing="(support for Guide §22)",
              purpose="A worked template for students who want a "
                      "starting point - not the only correct answer, "
                      "and not required.",
              board="None.",
              transition="Let a few students run their program for the "
                          "room.")
    return slide


def slide_19_activity_question(prs):
    slide = layout_big_question(
        prs, "Reflect",
        "The instant you reassign a\nvariable, what happens to its\nold "
        "value?", size=32,
        note="Completely replaced. Nothing is kept.")
    set_notes(slide, markers=["ASK", "PAUSE"], timing="(closes Guide "
              "§22)",
              purpose="Re-confirm Section 13's lesson using the "
                      "student's own program as the example.",
              board="None.",
              transition="Let's check a couple of misconceptions before "
                          "we close.")
    return slide


def slide_20_recap(prs):
    slide = new_slide(prs, bg=LIGHT_BG)
    y = _content_header(slide, "Quick Recap", "Quick Recap")
    draw_word_grid(["VARIABLE", "ASSIGNMENT", "INT", "FLOAT", "STR",
                     "BOOL"], cols=6, font_size=14)(
        slide, MARGIN, y + Inches(0.1), CONTENT_W, Inches(1.1))
    draw_pipeline(
        [("PROBLEM", None), ("LOGIC", None), ("ALGORITHM", None),
         ("PYTHON\nCODE", None), ("VARIABLES", None), ("RUN", None),
         ("OUTPUT/\nRESULT", None)],
        height=Inches(1.15), font_size=10.5, gap=Inches(0.24))(
        slide, MARGIN, y + Inches(1.7), CONTENT_W, Inches(1.4))
    add_footer(slide)
    set_notes(slide, markers=["ASK"], timing="2 min (Guide §24) — "
              "NON-NEGOTIABLE",
              purpose="One visual summary of every keyword and the full "
                      "bridged chain.",
              question="Why wasn't fixed text enough? What does = "
                        "actually do? What's the difference between "
                        "print(name) and print(\"name\")? What happens "
                        "to an old value after reassignment? Name the "
                        "four data types.",
              board="None.",
              transition="One question to leave you with.")
    return slide


def slide_21_bridge_forward(prs):
    slide = layout_big_question(
        prs, "A Question for Next Time",
        "We can print numbers now — but we\nhaven't done any math with "
        "them.\nWhat happens when a program needs\nto add, compare, or "
        "calculate?", size=27)
    set_notes(slide, markers=["DO NOT ANSWER"], timing="(Guide §25)",
              purpose="Leave this open as a deliberate hook into Class "
                      "07 (Operators & Expressions).",
              board="None.",
              transition="Class ends on the final takeaway.")
    return slide


def slide_22_final_takeaway(prs):
    slide = layout_section_divider(
        prs, "Final Takeaway",
        "“Last class, your programs could\nonly ever say the exact "
        "same thing.\nToday, they can hold something, change\nit, and "
        "tell you about it. That one idea\nis what makes a program "
        "feel alive\ninstead of frozen.”",
        title_size=26)
    set_notes(slide, timing="1 min",
              purpose="Close the class on this exact line, said slowly.",
              board="None.",
              transition="Class ends.")
    return slide


# ============================================================================
# 6. MAIN GENERATION FUNCTION
# ============================================================================

SLIDE_BUILDERS = [
    slide_01_title, slide_02_big_question, slide_03_callback,
    slide_04_what_is_a_variable, slide_05_assignment_statement,
    slide_06_assignment_proof, slide_07_print_name_vs_string,
    slide_08_reassignment, slide_09_naming_variables,
    slide_10_meet_data_types, slide_11_numbers, slide_12_strings_revisited,
    slide_13_booleans, slide_14_combining, slide_15_class04_callback,
    slide_16_full_bridge, slide_17_activity_intro,
    slide_18_activity_template, slide_19_activity_question,
    slide_20_recap, slide_21_bridge_forward, slide_22_final_takeaway,
]


def build_presentation(output_path):
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    for i, builder in enumerate(SLIDE_BUILDERS, start=1):
        slide = builder(prs)
        if i != 1:
            add_page_number(slide, i)
    prs.save(output_path)
    return len(SLIDE_BUILDERS)


if __name__ == "__main__":
    here = os.path.dirname(os.path.abspath(__file__))
    out = os.path.join(here, "Class_06_Variables_and_Data_Types.pptx")
    n = build_presentation(out)
    print(f"Generated {n} slides -> {out}")
