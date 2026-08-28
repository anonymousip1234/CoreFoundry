# -*- coding: utf-8 -*-
"""
generate_presentation.py

Generates Class_01_Welcome_to_Computer_Science.pptx for
Bong Study Hub - Foundation Batch 2026.

Source of truth (do not redesign the lesson when editing this script):
  01_Instructor_Guide/Class_01_Master_Instructor_Guide.md
  02_Student_Notes/Class_01_Student_Notes.md
  04_Pen_Tablet/Class_01_Pen_Tablet_Board_Plan.md
  04_Pen_Tablet/Class_01_Board_Quick_Reference.md
  05_Interaction/01_Question_Bank.md .. 06_Exit_Questions.md

This deck is the VISUAL SPINE of a live, instructor-led, pen-tablet-taught
class. It intentionally leaves space for live reasoning: several slides
are near-blank on purpose (see LIVE BOARD markers in speaker notes).

Structure of this file:
  1. Theme configuration (colors, fonts, spacing)
  2. Low-level utility functions (shapes, text, arrows)
  3. Reusable slide-layout components (title, question, diagram, ...)
  4. Content-aware diagram builders (pipeline, quadrant, roadmap cards...)
  5. Slide-specific builders (slide_01 .. slide_29)
  6. Main generation function
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR

# ============================================================================
# 1. THEME CONFIGURATION
# ============================================================================

# -- Colors ------------------------------------------------------------------
# Reused from the Class 01 Student Notes PDF palette for cross-artifact
# brand consistency across the whole program.
NAVY = RGBColor(0x1E, 0x3A, 0x72)          # primary — dark blue
BLUE = RGBColor(0x2A, 0x5C, 0xB8)          # secondary — blue
CYAN = RGBColor(0x22, 0xC7, 0xE0)          # accent — electric blue / cyan
INK = RGBColor(0x1F, 0x27, 0x33)           # body text
MUTED = RGBColor(0x64, 0x70, 0x85)         # secondary text
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_BG = RGBColor(0xF7, 0xF9, 0xFC)      # page background (light slides)
CARD_BG = RGBColor(0xEE, 0xF3, 0xFB)       # light navy-tinted card fill
CARD_BORDER = RGBColor(0xC7, 0xD6, 0xEC)
CYAN_BG = RGBColor(0xE4, 0xF9, 0xFB)
CYAN_BORDER = RGBColor(0x9F, 0xE6, 0xEE)
AMBER = RGBColor(0xC2, 0x79, 0x0F)         # rare — used only for the "tea
AMBER_BG = RGBColor(0xFB, 0xEE, 0xDC)      # failure" warning beat, matching
                                            # the Board Plan's color system.
RULE = RGBColor(0xDC, 0xE3, 0xED)
DOT_OFF = RGBColor(0xC7, 0xD1, 0xE0)

# -- Fonts ---------------------------------------------------------------
FONT_DISPLAY = "Aptos Display"   # large titles / statements
FONT_BODY = "Aptos"              # everything else
FONT_FALLBACK = "Arial"          # noted in README; Aptos ships with
                                  # current Windows/Office and degrades
                                  # gracefully if unavailable.

# -- Canvas ----------------------------------------------------------------
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
MARGIN = Inches(0.7)
CONTENT_W = SLIDE_W - 2 * MARGIN

# -- Program story stages (Visual Story arc — see PRESENTATION_README) ------
STORY_STAGES = [
    "Curiosity", "Computer Science", "How Computers Work", "Problem Solving",
    "Algorithms", "Computational Thinking", "Programming Precision", "AI",
    "Learning Journey", "Action",
]

FOOTER_TEXT = "Bong Study Hub  ·  Foundation Batch 2026  ·  Class 01"


# ============================================================================
# 2. LOW-LEVEL UTILITIES
# ============================================================================

def new_slide(prs, bg=LIGHT_BG):
    """Add a blank slide with a solid background fill."""
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = bg
    return slide


def _i(v):
    """Coerce any coordinate to an exact integer EMU. Required because
    Python's true division (e.g. `w / 2`) on a Length turns it into a
    plain float, and python-pptx writes that float straight into the
    XML as a non-integer string — which is invalid per the Open XML
    schema (ST_Coordinate requires an integer) and can make the shape
    unreadable. Every low-level shape-creation call below routes through
    this."""
    return int(round(v))


def _no_shadow(shape):
    try:
        shape.shadow.inherit = False
    except Exception:
        pass


def add_rect(slide, x, y, w, h, fill=WHITE, line=None, line_w=Pt(1),
             radius=0.06, shadow=False):
    """Rounded rectangle. radius=None gives a sharp-cornered rectangle."""
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shp = slide.shapes.add_shape(shape_type, _i(x), _i(y), _i(w), _i(h))
    if radius:
        try:
            shp.adjustments[0] = radius
        except Exception:
            pass
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = line_w
    if not shadow:
        _no_shadow(shp)
    shp.text_frame.margin_left = 0
    shp.text_frame.margin_right = 0
    shp.text_frame.margin_top = 0
    shp.text_frame.margin_bottom = 0
    return shp


def add_text(slide, x, y, w, h, text, size=18, color=INK, bold=False,
             italic=False, align=PP_ALIGN.LEFT, font=FONT_BODY,
             anchor=MSO_ANCHOR.TOP, line_spacing=1.08, wrap=True,
             shrink=False):
    """A textbox with one or more paragraphs (split on \\n)."""
    box = slide.shapes.add_textbox(_i(x), _i(y), _i(w), _i(h))
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    if shrink:
        tf.auto_size = MSO_AUTO_SIZE.NONE
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.name = font
        run.font.color.rgb = color
    return box


def add_rich_text(slide, x, y, w, h, segments, size=18, align=PP_ALIGN.LEFT,
                   font=FONT_BODY, anchor=MSO_ANCHOR.TOP, line_spacing=1.08):
    """One paragraph made of multiple (text, color, bold) runs."""
    box = slide.shapes.add_textbox(_i(x), _i(y), _i(w), _i(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    for text, color, bold in segments:
        run = p.add_run()
        run.text = text
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.name = font
        run.font.color.rgb = color
    return box


def add_line(slide, x1, y1, x2, y2, color=BLUE, width=Pt(1.5), dash=None):
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, _i(x1), _i(y1),
                                        _i(x2), _i(y2))
    conn.line.color.rgb = color
    conn.line.width = width
    _no_shadow(conn)
    if dash:
        from pptx.oxml.ns import qn
        ln = conn.line._get_or_add_ln()
        d = ln.makeelement(qn('a:prstDash'), {'val': dash})
        ln.append(d)
    return conn


def add_triangle(slide, cx, cy, size, color, rotation):
    """Small filled triangle used as an arrowhead. rotation: 90=right,
    180=down, 270=left, 0=up."""
    s = _i(size)
    shp = slide.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE,
                                  _i(cx) - s // 2, _i(cy) - s // 2, s, s)
    shp.rotation = rotation
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.line.fill.background()
    _no_shadow(shp)
    return shp


def add_arrow_h(slide, x1, y, x2, color=BLUE, width=Pt(1.5)):
    """Horizontal arrow from x1 to x2 at height y (pointing right)."""
    tip_gap = Emu(90000)
    add_line(slide, x1, y, x2 - tip_gap, y, color=color, width=width)
    add_triangle(slide, x2 - Emu(45000), y, Emu(160000), color, rotation=90)


def add_arrow_v(slide, x, y1, y2, color=BLUE, width=Pt(1.5)):
    """Vertical arrow from y1 to y2 at horizontal position x (pointing down)."""
    tip_gap = Emu(90000)
    add_line(slide, x, y1, x, y2 - tip_gap, color=color, width=width)
    add_triangle(slide, x, y2 - Emu(45000), Emu(160000), color, rotation=180)


def add_footer(slide):
    add_text(slide, MARGIN, SLIDE_H - Inches(0.42), Inches(6.5), Inches(0.3),
              FOOTER_TEXT, size=9.5, color=MUTED, font=FONT_BODY)


def add_page_number(slide, number):
    add_text(slide, SLIDE_W - Inches(1.1), SLIDE_H - Inches(0.42),
              Inches(0.5), Inches(0.3), str(number), size=9.5, color=MUTED,
              align=PP_ALIGN.RIGHT, font=FONT_BODY)


def add_progress_dots(slide, stage_index):
    """Subtle 10-dot progress row (Visual Story arc) near the footer."""
    n = len(STORY_STAGES)
    dot = Inches(0.09)
    gap = Inches(0.14)
    total_w = n * dot + (n - 1) * (gap - dot)
    start_x = SLIDE_W - MARGIN - total_w
    y = SLIDE_H - Inches(0.55)
    x = start_x
    for i in range(n):
        color = CYAN if i == stage_index else DOT_OFF
        shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, _i(x), _i(y), _i(dot), _i(dot))
        shp.fill.solid()
        shp.fill.fore_color.rgb = color
        shp.line.fill.background()
        _no_shadow(shp)
        x += gap


def add_kicker(slide, x, y, w, text, color=BLUE):
    add_text(slide, x, y, w, Inches(0.4), text.upper(), size=13, color=color,
              bold=True, font=FONT_BODY)


def set_notes(slide, purpose=None, say=None, question=None, expected=None,
               board=None, transition=None, timing=None, markers=None):
    """Speaker notes are reminders, not scripts. Markers appear only here,
    never on the visible slide."""
    lines = []
    if markers:
        lines.append("[" + "] [".join(markers) + "]")
    if timing:
        lines.append(f"TIME: {timing}")
    if purpose:
        lines.append(f"Purpose: {purpose}")
    if say:
        lines.append(f"Say: {say}")
    if question:
        lines.append(f"Ask: {question}")
    if expected:
        lines.append(f"Expected: {expected}")
    if board:
        lines.append(f"Board: {board}")
    if transition:
        lines.append(f"Transition: {transition}")
    slide.notes_slide.notes_text_frame.text = "\n".join(lines)


# ============================================================================
# 3. REUSABLE SLIDE-LAYOUT COMPONENTS
# ============================================================================

def layout_title(prs, kicker, title_lines, subtitle):
    """Layout 1 — TITLE. Full navy background, minimal, premium."""
    slide = new_slide(prs, bg=NAVY)
    # thin cyan accent rule
    add_rect(slide, MARGIN, Inches(2.15), Inches(1.2), Inches(0.06),
              fill=CYAN)
    add_kicker(slide, MARGIN, Inches(1.75), Inches(9), kicker, color=CYAN)
    y = Inches(2.4)
    for line in title_lines:
        add_text(slide, MARGIN, y, Inches(11.5), Inches(1.1), line, size=54,
                  color=WHITE, bold=True, font=FONT_DISPLAY)
        y += Inches(1.05)
    add_text(slide, MARGIN, y + Inches(0.15), Inches(10.5), Inches(0.6),
              subtitle, size=20, color=RGBColor(0xB9, 0xCC, 0xEE),
              italic=True, font=FONT_BODY)
    add_text(slide, MARGIN, SLIDE_H - Inches(0.9), Inches(6), Inches(0.5),
              "Bong Study Hub\nFoundation Batch 2026", size=13,
              color=RGBColor(0x8F, 0xA6, 0xD1), font=FONT_BODY,
              line_spacing=1.15)
    return slide


def layout_section_divider(prs, kicker, title, subtitle=None, stage_index=None):
    """Layout 2 — SECTION DIVIDER. Full navy background, one bold
    statement, used to mark a clear beat-change in the class story."""
    slide = new_slide(prs, bg=NAVY)
    add_rect(slide, MARGIN, Inches(3.05), Inches(1.0), Inches(0.06), fill=CYAN)
    if kicker:
        add_kicker(slide, MARGIN, Inches(2.65), Inches(10), kicker, color=CYAN)
    add_text(slide, MARGIN, Inches(3.25), Inches(11.8), Inches(2.0), title,
              size=44, color=WHITE, bold=True, font=FONT_DISPLAY,
              line_spacing=1.05)
    if subtitle:
        add_text(slide, MARGIN, Inches(5.05), Inches(10.5), Inches(0.8),
                  subtitle, size=18, color=RGBColor(0xB9, 0xCC, 0xEE),
                  font=FONT_BODY, italic=True)
    add_text(slide, MARGIN, SLIDE_H - Inches(0.42), Inches(6.5), Inches(0.3),
              FOOTER_TEXT, size=9.5, color=RGBColor(0x6B, 0x82, 0xAE))
    if stage_index is not None:
        add_progress_dots(slide, stage_index)
    return slide


def _content_header(slide, kicker, title, title_size=32):
    if kicker:
        add_kicker(slide, MARGIN, Inches(0.5), CONTENT_W, kicker, color=BLUE)
        title_y = Inches(0.86)
    else:
        title_y = Inches(0.6)
    add_text(slide, MARGIN, title_y, CONTENT_W, Inches(0.9), title,
              size=title_size, color=NAVY, bold=True, font=FONT_DISPLAY)
    add_rect(slide, MARGIN, title_y + Inches(0.85), Inches(0.85),
              Inches(0.045), fill=CYAN)
    return title_y + Inches(1.15)  # y where content area starts


def layout_big_question(prs, kicker, question, note=None, stage_index=None):
    """Layout 3 — BIG QUESTION. Minimal, large centered text, no answer."""
    slide = new_slide(prs, bg=LIGHT_BG)
    if kicker:
        add_kicker(slide, MARGIN, Inches(1.5), CONTENT_W, kicker,
                    color=BLUE)
    add_text(slide, Inches(1.1), Inches(2.5), Inches(11.1), Inches(2.6),
              question, size=40, color=NAVY, bold=True, font=FONT_DISPLAY,
              align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
              line_spacing=1.12)
    if note:
        add_text(slide, Inches(1.1), Inches(5.3), Inches(11.1), Inches(0.5),
                  note, size=15, color=MUTED, align=PP_ALIGN.CENTER,
                  italic=True)
    add_footer(slide)
    if stage_index is not None:
        add_progress_dots(slide, stage_index)
    return slide


def layout_concept_explanation(prs, kicker, title, statement=None,
                                bullets=None, stage_index=None,
                                title_size=32):
    """Layout 4 — CONCEPT EXPLANATION. Short title + either one big
    statement or a small set of short bullets — never a paragraph."""
    slide = new_slide(prs, bg=LIGHT_BG)
    y = _content_header(slide, kicker, title, title_size=title_size)
    if statement:
        add_text(slide, MARGIN, y + Inches(0.3), CONTENT_W, Inches(2.2),
                  statement, size=26, color=INK, font=FONT_BODY,
                  line_spacing=1.2)
    if bullets:
        yy = y + Inches(0.35)
        for b in bullets:
            add_rect(slide, MARGIN, yy + Inches(0.14), Inches(0.14),
                      Inches(0.14), fill=CYAN, radius=0.5)
            add_text(slide, MARGIN + Inches(0.35), yy, CONTENT_W - Inches(0.35),
                      Inches(0.55), b, size=19, color=INK)
            yy += Inches(0.62)
    add_footer(slide)
    if stage_index is not None:
        add_progress_dots(slide, stage_index)
    return slide


def layout_two_column(prs, kicker, title, left, right, connector_label=None,
                       stage_index=None):
    """Layout 5 — TWO-COLUMN COMPARISON.
    left/right = {"title": str, "sub": str, "tag": str}
    """
    slide = new_slide(prs, bg=LIGHT_BG)
    y = _content_header(slide, kicker, title)
    col_w = (CONTENT_W - Inches(0.6)) / 2
    card_h = Inches(2.7)
    card_y = y + Inches(0.4)
    for i, card in enumerate((left, right)):
        x = MARGIN + i * (col_w + Inches(0.6))
        fill = CARD_BG if i == 0 else CYAN_BG
        border = CARD_BORDER if i == 0 else CYAN_BORDER
        add_rect(slide, x, card_y, col_w, card_h, fill=fill, line=border,
                  line_w=Pt(1.25), radius=0.05)
        if card.get("tag"):
            add_text(slide, x + Inches(0.35), card_y + Inches(0.3),
                      col_w - Inches(0.7), Inches(0.4), card["tag"],
                      size=13, color=(BLUE if i == 0 else NAVY), bold=True)
        add_text(slide, x + Inches(0.35), card_y + Inches(0.75),
                  col_w - Inches(0.7), Inches(0.7), card["title"], size=27,
                  color=NAVY, bold=True, font=FONT_DISPLAY)
        add_text(slide, x + Inches(0.35), card_y + Inches(1.55),
                  col_w - Inches(0.7), Inches(1.0), card["sub"], size=17,
                  color=INK, line_spacing=1.2)
    # connector arrow between the two cards
    mid_y = card_y + card_h / 2
    gap_x1 = MARGIN + col_w
    gap_x2 = gap_x1 + Inches(0.6)
    add_arrow_h(slide, gap_x1 + Inches(0.08), mid_y, gap_x2 - Inches(0.08),
                color=BLUE, width=Pt(2))
    if connector_label:
        add_text(slide, gap_x1, mid_y - Inches(0.55), Inches(0.6),
                  Inches(0.35), connector_label, size=11, color=MUTED,
                  align=PP_ALIGN.CENTER)
    add_footer(slide)
    if stage_index is not None:
        add_progress_dots(slide, stage_index)
    return slide


def layout_diagram(prs, kicker, title, diagram_fn, caption=None,
                    stage_index=None, title_size=32):
    """Layout 6 — DIAGRAM. Title + a content-aware diagram drawn by
    diagram_fn(slide, x, y, w, h)."""
    slide = new_slide(prs, bg=LIGHT_BG)
    y = _content_header(slide, kicker, title, title_size=title_size)
    area_y = y + Inches(0.25)
    area_h = SLIDE_H - area_y - Inches(0.75)
    if caption:
        area_h -= Inches(0.5)
    diagram_fn(slide, MARGIN, area_y, CONTENT_W, area_h)
    if caption:
        add_text(slide, MARGIN, SLIDE_H - Inches(1.05), CONTENT_W,
                  Inches(0.4), caption, size=13, color=MUTED, italic=True,
                  align=PP_ALIGN.CENTER)
    add_footer(slide)
    if stage_index is not None:
        add_progress_dots(slide, stage_index)
    return slide


def layout_full_screen_visual(prs, kicker, title, diagram_fn,
                               stage_index=None):
    """Layout 7 — FULL-SCREEN VISUAL. Minimal chrome, diagram dominates."""
    slide = new_slide(prs, bg=LIGHT_BG)
    add_kicker(slide, MARGIN, Inches(0.45), CONTENT_W, kicker, color=BLUE)
    add_text(slide, MARGIN, Inches(0.8), CONTENT_W, Inches(0.7), title,
              size=28, color=NAVY, bold=True, font=FONT_DISPLAY)
    diagram_fn(slide, MARGIN, Inches(1.75), CONTENT_W,
                SLIDE_H - Inches(1.75) - Inches(1.3))
    add_text(slide, MARGIN, SLIDE_H - Inches(1.05), CONTENT_W, Inches(0.4),
              "A simplified mental model — not the real architecture.",
              size=13, color=MUTED, italic=True, align=PP_ALIGN.CENTER)
    add_footer(slide)
    if stage_index is not None:
        add_progress_dots(slide, stage_index)
    return slide


def layout_activity(prs, kicker, title, prompt_lines, sub=None, tag="ACTIVITY",
                     stage_index=None):
    """Layout 8 — ACTIVITY. A clear prompt card signalling hands-on work."""
    slide = new_slide(prs, bg=LIGHT_BG)
    add_rect(slide, Inches(0.9), Inches(0.55), Inches(1.5), Inches(0.42),
              fill=CYAN_BG, line=CYAN_BORDER, line_w=Pt(1))
    add_text(slide, Inches(0.9), Inches(0.55), Inches(1.5), Inches(0.42),
              tag, size=13, color=NAVY, bold=True, align=PP_ALIGN.CENTER,
              anchor=MSO_ANCHOR.MIDDLE)
    add_kicker(slide, Inches(2.6), Inches(0.62), Inches(9), kicker,
                color=BLUE)
    add_text(slide, Inches(0.9), Inches(1.25), Inches(11.5), Inches(1.0),
              title, size=32, color=NAVY, bold=True, font=FONT_DISPLAY)
    y = Inches(2.55)
    for line in prompt_lines:
        add_text(slide, Inches(0.9), y, Inches(11.5), Inches(0.85), line,
                  size=26, color=INK, font=FONT_BODY, line_spacing=1.15)
        y += Inches(0.85)
    if sub:
        add_text(slide, Inches(0.9), y + Inches(0.15), Inches(11.0),
                  Inches(0.6), sub, size=16, color=MUTED, italic=True)
    add_footer(slide)
    if stage_index is not None:
        add_progress_dots(slide, stage_index)
    return slide


def layout_roadmap(prs, kicker, title, cards, stage_index=None):
    """Layout 9 — ROADMAP. Four (or n) horizontal cards."""
    slide = new_slide(prs, bg=LIGHT_BG)
    y = _content_header(slide, kicker, title)
    n = len(cards)
    gap = Inches(0.35)
    card_w = (CONTENT_W - (n - 1) * gap) / n
    card_y = y + Inches(0.35)
    card_h = Inches(4.35)
    for i, card in enumerate(cards):
        x = MARGIN + i * (card_w + gap)
        add_rect(slide, x, card_y, card_w, card_h, fill=WHITE,
                  line=CARD_BORDER, line_w=Pt(1), radius=0.04)
        add_rect(slide, x, card_y, card_w, Inches(0.55), fill=NAVY,
                  radius=0.04)
        add_rect(slide, x, card_y + Inches(0.28), card_w, Inches(0.27),
                  fill=NAVY, radius=None)  # square off bottom of header
        add_text(slide, x, card_y, card_w, Inches(0.55), card["month"],
                  size=15, color=WHITE, bold=True, align=PP_ALIGN.CENTER,
                  anchor=MSO_ANCHOR.MIDDLE)
        add_text(slide, x + Inches(0.25), card_y + Inches(0.75),
                  card_w - Inches(0.5), Inches(0.6), card["theme"], size=17,
                  color=NAVY, bold=True, font=FONT_BODY, line_spacing=1.05)
        yy = card_y + Inches(1.5)
        for item in card["items"]:
            add_text(slide, x + Inches(0.28), yy, card_w - Inches(0.5),
                      Inches(0.4), "• " + item, size=13.5, color=INK)
            yy += Inches(0.42)
    add_footer(slide)
    if stage_index is not None:
        add_progress_dots(slide, stage_index)
    return slide


def layout_recap(prs, kicker, title, items, numbered=False, stage_index=None):
    """Layout 10 — RECAP. Short list of statements or questions."""
    slide = new_slide(prs, bg=LIGHT_BG)
    y = _content_header(slide, kicker, title)
    yy = y + Inches(0.3)
    row_h = (SLIDE_H - Inches(0.9) - yy) / len(items)
    row_h = min(row_h, Inches(0.62))
    for i, item in enumerate(items):
        if numbered:
            add_rect(slide, MARGIN, yy + Inches(0.03), Inches(0.4),
                      Inches(0.4), fill=NAVY, radius=0.5)
            add_text(slide, MARGIN, yy + Inches(0.03), Inches(0.4),
                      Inches(0.4), str(i + 1), size=15, color=WHITE,
                      bold=True, align=PP_ALIGN.CENTER,
                      anchor=MSO_ANCHOR.MIDDLE)
            text_x = MARGIN + Inches(0.58)
        else:
            add_rect(slide, MARGIN, yy + Inches(0.16), Inches(0.14),
                      Inches(0.14), fill=CYAN, radius=0.5)
            text_x = MARGIN + Inches(0.4)
        add_text(slide, text_x, yy, CONTENT_W - Inches(0.6), row_h, item,
                  size=18, color=INK, anchor=MSO_ANCHOR.TOP,
                  line_spacing=1.1)
        yy += row_h
    add_footer(slide)
    if stage_index is not None:
        add_progress_dots(slide, stage_index)
    return slide


def layout_homework(prs, kicker, title, subtitle, list_items, steps,
                     stage_index=None):
    """Layout 11 — HOMEWORK. Clean, minimal closing slide."""
    slide = new_slide(prs, bg=NAVY)
    add_kicker(slide, MARGIN, Inches(0.6), CONTENT_W, kicker, color=CYAN)
    add_text(slide, MARGIN, Inches(0.95), CONTENT_W, Inches(0.9), title,
              size=38, color=WHITE, bold=True, font=FONT_DISPLAY)
    add_text(slide, MARGIN, Inches(1.75), CONTENT_W, Inches(0.5), subtitle,
              size=17, color=RGBColor(0xB9, 0xCC, 0xEE), italic=True)

    col_w = Inches(5.6)
    # left: suggested activities
    add_text(slide, MARGIN, Inches(2.55), col_w, Inches(0.4),
              "CHOOSE ONE EVERYDAY PROCESS", size=13, color=CYAN, bold=True)
    yy = Inches(3.0)
    cols = 2
    item_w = col_w / cols
    for i, item in enumerate(list_items):
        cx = MARGIN + (i % cols) * item_w
        cy = yy + (i // cols) * Inches(0.45)
        add_text(slide, cx, cy, item_w - Inches(0.1), Inches(0.4),
                  "• " + item, size=15, color=WHITE)

    # right: steps
    right_x = MARGIN + col_w + Inches(0.5)
    add_text(slide, right_x, Inches(2.55), Inches(5.8), Inches(0.4),
              "FOR YOUR CHOSEN PROCESS", size=13, color=CYAN, bold=True)
    yy = Inches(3.0)
    for i, step in enumerate(steps):
        add_rect(slide, right_x, yy + Inches(0.02), Inches(0.36),
                  Inches(0.36), fill=CYAN, radius=0.5)
        add_text(slide, right_x, yy + Inches(0.02), Inches(0.36),
                  Inches(0.36), str(i + 1), size=14, color=NAVY, bold=True,
                  align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        add_text(slide, right_x + Inches(0.5), yy, Inches(5.3), Inches(0.6),
                  step, size=16, color=WHITE, line_spacing=1.1)
        yy += Inches(0.72)

    add_text(slide, MARGIN, SLIDE_H - Inches(0.42), Inches(6.5),
              Inches(0.3), FOOTER_TEXT, size=9.5,
              color=RGBColor(0x6B, 0x82, 0xAE))
    if stage_index is not None:
        add_progress_dots(slide, stage_index)
    return slide


# ============================================================================
# 4. CONTENT-AWARE DIAGRAM BUILDERS
# ============================================================================

def draw_pipeline(items, height=Inches(1.3), font_size=15, highlight_last=False):
    """Returns a diagram_fn(slide, x, y, w, h) drawing a horizontal chain
    of rounded boxes connected by arrows. items: list of (label, sub)."""
    def _fn(slide, x, y, w, h):
        n = len(items)
        gap = Inches(0.55)
        box_w = (w - gap * (n - 1)) / n
        box_y = y + (h - height) / 2
        cx = x
        for i, (label, sub) in enumerate(items):
            is_last = highlight_last and i == n - 1
            fill = CYAN_BG if is_last else CARD_BG
            border = CYAN_BORDER if is_last else CARD_BORDER
            add_rect(slide, cx, box_y, box_w, height, fill=fill,
                      line=border, line_w=Pt(1.25), radius=0.09)
            label_y = box_y + (height / 2 - Inches(0.32) if sub
                                else height / 2 - Inches(0.24))
            add_text(slide, cx, label_y, box_w, Inches(0.5), label,
                      size=font_size, color=NAVY, bold=True,
                      align=PP_ALIGN.CENTER, font=FONT_BODY)
            if sub:
                add_text(slide, cx, label_y + Inches(0.42), box_w,
                          Inches(0.35), sub, size=font_size - 4, color=MUTED,
                          align=PP_ALIGN.CENTER)
            if i < n - 1:
                add_arrow_h(slide, cx + box_w + Inches(0.06), box_y + height / 2,
                            cx + box_w + gap - Inches(0.06), color=BLUE,
                            width=Pt(1.75))
            cx += box_w + gap
    return _fn


def draw_stacked_pipelines(rows, row_h=Inches(1.05), gap=Inches(0.3),
                            font_size=13):
    """rows: list of (row_label, [(label, None), ...], highlight_idx)"""
    def _fn(slide, x, y, w, h):
        label_w = Inches(2.15)
        chain_x = x + label_w
        chain_w = w - label_w
        cy = y
        for row_label, items, highlight_idx in rows:
            add_text(slide, x, cy, label_w - Inches(0.2), row_h, row_label,
                      size=font_size + 1, color=NAVY, bold=True,
                      anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.05)
            n = len(items)
            box_gap = Inches(0.4)
            box_w = (chain_w - box_gap * (n - 1)) / n
            cx = chain_x
            for i, label in enumerate(items):
                is_hl = (i == highlight_idx)
                fill = CYAN_BG if is_hl else CARD_BG
                border = CYAN_BORDER if is_hl else CARD_BORDER
                add_rect(slide, cx, cy, box_w, row_h, fill=fill,
                          line=border, line_w=Pt(1), radius=0.12)
                add_text(slide, cx, cy, box_w, row_h, label,
                          size=font_size, color=NAVY, bold=True,
                          align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
                if i < n - 1:
                    add_arrow_h(slide, cx + box_w + Inches(0.04), cy + row_h / 2,
                                cx + box_w + box_gap - Inches(0.04),
                                color=BLUE, width=Pt(1.4))
                cx += box_w + box_gap
            cy += row_h + gap
    return _fn


CONNECTOR_LINE = RGBColor(0xA7, 0xBC, 0xE3)  # visible but subtle — more
                                              # contrast than CARD_BORDER


def draw_cs_map(areas):
    """Height-responsive: card size and gaps scale to whatever vertical
    space the caller provides, so this never overflows into a caption or
    footer regardless of which layout hosts it."""
    def _fn(slide, x, y, w, h):
        cols = 3
        rows = 3
        center_w = min(Inches(3.6), w * 0.42)
        center_h = Inches(0.75)
        cx = x + w / 2 - center_w / 2
        cy = y
        top_gap = Inches(0.4)
        row_gap = Inches(0.22)
        card_gap_x = Inches(0.3)

        available_h = (y + h) - (cy + center_h + top_gap)
        card_h = (available_h - row_gap * (rows - 1)) / rows
        card_h = max(Inches(0.55), min(card_h, Inches(0.85)))
        grid_top = cy + center_h + top_gap

        add_rect(slide, cx, cy, center_w, center_h, fill=NAVY, radius=0.18)
        add_text(slide, cx, cy, center_w, center_h, "COMPUTER SCIENCE",
                  size=18, color=WHITE, bold=True, align=PP_ALIGN.CENTER,
                  anchor=MSO_ANCHOR.MIDDLE)

        # Org-chart-style elbow connectors (trunk + column branches)
        # instead of direct spokes — direct lines from one apex to nine
        # scattered cards cross visually through unrelated cards, which
        # reads as cluttered rather than premium.
        card_w = (w - card_gap_x * (cols - 1)) / cols
        col_centers = [x + c * (card_w + card_gap_x) + card_w / 2
                        for c in range(cols)]
        center_x = x + w / 2
        trunk_y = cy + center_h + top_gap / 2
        add_line(slide, center_x, cy + center_h, center_x, trunk_y,
                  color=CONNECTOR_LINE, width=Pt(1.25))
        add_line(slide, col_centers[0], trunk_y, col_centers[-1], trunk_y,
                  color=CONNECTOR_LINE, width=Pt(1.25))
        for cxpos in col_centers:
            add_line(slide, cxpos, trunk_y, cxpos, grid_top,
                      color=CONNECTOR_LINE, width=Pt(1.25))

        for i, area in enumerate(areas):
            r, c = divmod(i, cols)
            bx = x + c * (card_w + card_gap_x)
            by = grid_top + r * (card_h + row_gap)
            if r > 0:
                # short vertical link from the card above, same column
                prev_bottom = by - row_gap
                add_line(slide, col_centers[c], prev_bottom, col_centers[c],
                          by, color=CONNECTOR_LINE, width=Pt(1.25))
            add_rect(slide, bx, by, card_w, card_h, fill=CARD_BG,
                      line=CARD_BORDER, line_w=Pt(1), radius=0.12)
            add_text(slide, bx + Inches(0.1), by, card_w - Inches(0.2),
                      card_h, area, size=14, color=NAVY, bold=False,
                      align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
                      line_spacing=1.0)
    return _fn


def draw_vertical_flow(items):
    def _fn(slide, x, y, w, h):
        n = len(items)
        box_w = Inches(4.6)
        box_h = Inches(0.62)
        gap = (h - n * box_h) / (n - 1) if n > 1 else Inches(0)
        gap = min(gap, Inches(0.45))
        total_h = n * box_h + (n - 1) * gap
        cy = y + (h - total_h) / 2
        bx = x + w / 2 - box_w / 2
        for i, label in enumerate(items):
            add_rect(slide, bx, cy, box_w, box_h, fill=CARD_BG,
                      line=CARD_BORDER, line_w=Pt(1), radius=0.14)
            add_text(slide, bx, cy, box_w, box_h, label, size=16,
                      color=NAVY, bold=True, align=PP_ALIGN.CENTER,
                      anchor=MSO_ANCHOR.MIDDLE)
            if i < n - 1:
                add_arrow_v(slide, x + w / 2, cy + box_h + Inches(0.04),
                            cy + box_h + gap - Inches(0.04), color=BLUE,
                            width=Pt(1.5))
            cy += box_h + gap
    return _fn


def draw_quadrants(items):
    """items: list of 4 dicts {title, sub}. Grid height is capped and
    vertically centered rather than stretching to fill the full area —
    these are short word+example cards, not paragraphs."""
    def _fn(slide, x, y, w, h):
        gap = Inches(0.35)
        cw = (w - gap) / 2
        ch = min(Inches(1.85), (h - gap) / 2)
        grid_h = ch * 2 + gap
        top = y + max(Inches(0), (h - grid_h) / 2)
        positions = [(x, top), (x + cw + gap, top), (x, top + ch + gap),
                     (x + cw + gap, top + ch + gap)]
        fills = [CARD_BG, CYAN_BG, CYAN_BG, CARD_BG]
        borders = [CARD_BORDER, CYAN_BORDER, CYAN_BORDER, CARD_BORDER]
        for (px, py), item, fill, border in zip(positions, items, fills, borders):
            add_rect(slide, px, py, cw, ch, fill=fill, line=border,
                      line_w=Pt(1.25), radius=0.08)
            add_text(slide, px + Inches(0.3), py + Inches(0.26),
                      cw - Inches(0.6), Inches(0.5), item["title"], size=18,
                      color=NAVY, bold=True, font=FONT_BODY)
            add_text(slide, px + Inches(0.3), py + Inches(0.75),
                      cw - Inches(0.6), ch - Inches(0.95), item["sub"],
                      size=14, color=INK, line_spacing=1.15)
    return _fn


def draw_staircase(labels, highlight_idx=(0, 1)):
    """Reserves a fixed label zone (two lines) at the bottom of h so
    two-line labels like "Machine\\nLearning" never extend past the
    area the caller allocated for this diagram."""
    def _fn(slide, x, y, w, h):
        n = len(labels)
        label_zone = Inches(0.65)
        bar_area_h = h - label_zone
        step_w = (w - Inches(0.15) * (n - 1)) / n
        base_h = bar_area_h * 0.20
        top_h = bar_area_h * 0.62
        baseline_y = y + bar_area_h
        cx = x
        for i in range(n):
            step_h = base_h + (top_h - base_h) * (i / (n - 1))
            step_y = baseline_y - step_h
            is_hl = i in highlight_idx
            fill = CYAN_BG if is_hl else CARD_BG
            border = CYAN if is_hl else CARD_BORDER
            add_rect(slide, cx, step_y, step_w, step_h, fill=fill,
                      line=border, line_w=Pt(1.5 if is_hl else 1),
                      radius=0.1)
            add_text(slide, cx, step_y - Inches(0.42), step_w, Inches(0.35),
                      str(i + 1), size=15, color=(NAVY if is_hl else MUTED),
                      bold=True, align=PP_ALIGN.CENTER)
            cx += step_w + Inches(0.15)
        # labels below baseline
        label_y = baseline_y + Inches(0.12)
        cx = x
        for i, label in enumerate(labels):
            is_hl = i in highlight_idx
            add_text(slide, cx - Inches(0.1), label_y, step_w + Inches(0.2),
                      label_zone - Inches(0.12), label, size=11.5,
                      color=(NAVY if is_hl else MUTED),
                      bold=is_hl, align=PP_ALIGN.CENTER, line_spacing=1.0)
            cx += step_w + Inches(0.15)
    return _fn


def draw_three_words_cards(cards):
    """cards: list of 3 dicts {title, sub}. Card height is fixed to fit
    its content (title + 2-line definition) rather than stretching to
    fill whatever area it's given — avoids large empty cards."""
    def _fn(slide, x, y, w, h):
        gap = Inches(0.4)
        n = len(cards)
        card_w = (w - gap * (n - 1)) / n
        card_h = min(Inches(2.5), h)
        top = y + max(Inches(0), (h - card_h) / 2)
        for i, card in enumerate(cards):
            cx = x + i * (card_w + gap)
            add_rect(slide, cx, top, card_w, card_h, fill=WHITE,
                      line=CARD_BORDER, line_w=Pt(1), radius=0.07)
            add_rect(slide, cx, top, card_w, Inches(0.09), fill=CYAN,
                      radius=None)
            add_text(slide, cx + Inches(0.3), top + Inches(0.4),
                      card_w - Inches(0.6), Inches(0.7), card["title"],
                      size=24, color=NAVY, bold=True, font=FONT_DISPLAY)
            add_text(slide, cx + Inches(0.3), top + Inches(1.2),
                      card_w - Inches(0.6), card_h - Inches(1.5),
                      card["sub"], size=15, color=INK, line_spacing=1.25)
    return _fn


def draw_mcq(question, options):
    def _fn(slide, x, y, w, h):
        add_text(slide, x, y, w, Inches(1.0), question, size=30, color=NAVY,
                  bold=True, font=FONT_DISPLAY)
        oy = y + Inches(1.25)
        letters = ["A", "B", "C", "D"]
        for letter, opt in zip(letters, options):
            add_rect(slide, x, oy, Inches(0.55), Inches(0.55), fill=CARD_BG,
                      line=CARD_BORDER, line_w=Pt(1), radius=0.5)
            add_text(slide, x, oy, Inches(0.55), Inches(0.55), letter,
                      size=18, color=NAVY, bold=True, align=PP_ALIGN.CENTER,
                      anchor=MSO_ANCHOR.MIDDLE)
            add_text(slide, x + Inches(0.75), oy, w - Inches(0.75),
                      Inches(0.55), opt, size=19, color=INK,
                      anchor=MSO_ANCHOR.MIDDLE)
            oy += Inches(0.75)
    return _fn


def draw_single_word(word, size=90, color=NAVY):
    def _fn(slide, x, y, w, h):
        add_text(slide, x, y, w, h, word, size=size, color=color, bold=True,
                  font=FONT_DISPLAY, align=PP_ALIGN.LEFT,
                  anchor=MSO_ANCHOR.MIDDLE)
    return _fn


# ============================================================================
# 5. SLIDE-SPECIFIC BUILDERS
# ============================================================================

def slide_01_title(prs):
    slide = layout_title(
        prs, "Bong Study Hub · Foundation Batch 2026 · Class 01",
        ["Welcome to", "Computer Science"],
        "From Using Technology → Understanding Technology")
    set_notes(slide, timing="2 min",
              purpose="Set tone: welcoming, premium, unhurried.",
              say="Introduce yourself briefly; do not read the slide aloud.",
              board="None.",
              transition="How many of you used a smartphone today?")
    return slide


def slide_02_opening_question(prs):
    slide = layout_big_question(
        prs, "Opening", "How much of the technology you use every day\ndo you actually understand?",
        note=None, stage_index=0)
    set_notes(slide, markers=["ASK", "PAUSE"], timing="3 min",
              purpose="Start conversation; do not answer this on-slide.",
              question="How many of you used a smartphone / a computer today? "
                        "Who knows what happens inside when you press a button?",
              expected="Broad hand-raise participation, not verbal answers yet.",
              board="None.",
              transition="We're going to start understanding that today. "
                          "What is Computer Science?")
    return slide


def slide_03_what_is_cs(prs):
    slide = layout_big_question(prs, "Discussion",
                                  "What is\nComputer Science?", stage_index=1)
    set_notes(slide, markers=["ASK", "LET STUDENTS ANSWER", "DISCUSS"],
              timing="5 min",
              purpose="Surface students' existing mental model before "
                      "correcting it (Question Bank Q1.2).",
              question="What is Computer Science?",
              expected="Coding, programming, computers, AI, software — "
                        "collect 3–5 answers without judgment.",
              board="None yet — answers feed the live CS map next.",
              transition="Let's use exactly what you just said to build a map.")
    return slide


def slide_04_cs_not_coding(prs):
    slide = layout_two_column(
        prs, "The Reframe", "Computer Science ≠ Coding",
        {"tag": "ONE PART OF THE FIELD", "title": "Coding",
         "sub": "Writing instructions in a programming language. "
                "Visible, popular — and only one skill inside CS."},
        {"tag": "THE WHOLE FIELD", "title": "Computer Science",
         "sub": "Computation, information, algorithms, and problem "
                "solving — coding is one room in a much bigger "
                "building."},
        connector_label="part of",
        stage_index=1)
    set_notes(slide, markers=["DISCUSS"], timing="2 min",
              purpose="State the reframe plainly after the opening "
                      "discussion (Guide §5.2).",
              say="You were all partly right — you each named one room "
                  "in a bigger building.",
              board="None.",
              transition="Let's see the whole building.")
    return slide


def slide_05_cs_map(prs):
    areas = ["Programming", "Algorithms", "Data Structures",
              "Operating Systems", "Databases", "Computer Networks",
              "Computer Architecture", "Security", "Artificial Intelligence"]
    slide = layout_diagram(prs, "Reference", "The Computer Science Map",
                             draw_cs_map(areas),
                             caption="Clean reference — built live on "
                                     "the board first.",
                             stage_index=1, title_size=30)
    set_notes(slide, markers=["LIVE BOARD"], timing="8 min (Board Plan Drawing 1, ~9–17 min)",
              purpose="This is the CLEAN REFERENCE version only. The "
                      "instructor draws this map live on the pen tablet, "
                      "built from students' own opening answers — do "
                      "not present this slide before the live drawing.",
              say="You said coding, AI, software — all of you were "
                  "right, you each named one room in a bigger building.",
              board="LIVE BOARD — instructor draws the CS map "
                    "(Board Plan Drawing 1). Show this slide only after "
                    "the live drawing is complete, as a clean callback.",
              transition="Now that we've mapped the field, let's zoom into "
                          "one question: how does a computer actually "
                          "solve something?")
    return slide


def slide_06_quick_question(prs):
    slide = layout_big_question(
        prs, "Quick Question",
        "What part of Computer Science\nhave you already touched?",
        note="Apps · Games · Internet · AI · Databases · Security",
        stage_index=1)
    set_notes(slide, markers=["ASK", "LET STUDENTS ANSWER"], timing="2 min",
              purpose="Make the CS map personally relevant (Question Bank Q2.1).",
              question="Which of these have you touched without realizing "
                        "it was Computer Science?",
              expected="Using an app, playing a game, connecting to Wi-Fi.",
              board="Map stays visible if convenient.",
              transition="So which one of these boxes is programming?")
    return slide


def slide_07_button_press(prs):
    flow = ["User", "Device", "Application", "Network", "Server",
            "Database", "Response", "Device"]
    slide = layout_full_screen_visual(
        prs, "The Hook", "What Happens When You Press a Button?",
        draw_vertical_flow(flow), stage_index=2)
    set_notes(slide, markers=["DISCUSS"], timing="3 min (optional — cut "
              "first under time pressure, Guide §4.3)",
              purpose="Show a 'simple' action hides real computing "
                      "concepts — motivates why CS is worth studying "
                      "(Guide §5.3).",
              say="This is a simplified conceptual model, not the real "
                  "architecture — we are not doing networking today.",
              board="None (optional quick trace on board if energy is "
                    "high).",
              transition="Let's get precise vocabulary for three of "
                          "those words.")
    return slide


def slide_08_three_words(prs):
    cards = [
        {"title": "Computer", "sub": "A machine capable of executing "
                                       "instructions."},
        {"title": "Program", "sub": "A sequence of instructions a "
                                      "computer can execute."},
        {"title": "Application", "sub": "Software designed to perform a "
                                          "useful task for users."},
    ]
    slide = layout_diagram(prs, "Vocabulary", "Three Words",
                             draw_three_words_cards(cards),
                             stage_index=2, title_size=32)
    set_notes(slide, timing="3 min",
              purpose="Give precise, reusable vocabulary (Guide §5.4).",
              say="Examples: calculator, browser, game, messaging app, AI "
                  "chatbot — for each, ask 'program, application, or "
                  "both?'",
              board="None — this block is intentionally a slide "
                    "reference, not a live drawing.",
              transition="Let's test that with a question.")
    return slide


def slide_09_quick_check(prs):
    slide = layout_diagram(
        prs, "Quick Check", "",
        draw_mcq("A calculator app is a...",
                  ["Program", "Application", "Both", "Neither"]),
        stage_index=2, title_size=1)
    # remove the empty title header artifact by covering with bg rect is
    # unnecessary since title_size=1 keeps it invisible-small; kicker still
    # shows "Quick Check" as the anchor label.
    set_notes(slide, markers=["ASK", "LET STUDENTS ANSWER"], timing="2 min",
              purpose="Test program vs. application distinction "
                      "(Question Bank Q3.1).",
              question="A calculator app is a program, an application, "
                        "or both?",
              expected="'Both' — confirm and explain why: program is "
                        "the instructions, application is the packaged, "
                        "useful version for a user.",
              board="None.",
              transition="That 'Process' box was a mystery box. Let's "
                          "open the idea of Input → Process → Output.")
    return slide


def slide_10_ipo(prs):
    items = [("INPUT", "15 + 20"), ("PROCESS", "Addition"), ("OUTPUT", "35")]
    slide = layout_diagram(prs, "Reference", "Input → Process → Output",
                             draw_pipeline(items, height=Inches(1.6),
                                            font_size=20),
                             caption="Clean reference — built live on "
                                     "the board first.",
                             stage_index=2)
    set_notes(slide, markers=["LIVE BOARD"],
              timing="6 min (Board Plan Drawing 2, ~33–39 min)",
              purpose="Simple mental model for how computers solve "
                      "problems (Guide §5.5). Clean reference only — "
                      "instructor builds this live, box by box.",
              board="LIVE BOARD — instructor draws Drawing 2. Show "
                    "this slide only after the live drawing.",
              say="Process stays a black box today — on purpose. Real "
                  "systems can be far more complex than three boxes.",
              transition="Now find your own example, in pairs.")
    return slide


def slide_11_your_turn(prs):
    slide = layout_activity(
        prs, "Pair Activity", "Your Turn",
        ["Pick something you did today.", "What was the INPUT? PROCESS? OUTPUT?"],
        sub="90 seconds with the person next to you.",
        stage_index=2)
    set_notes(slide, markers=["ACTIVITY", "PAUSE"], timing="4 min "
              "(Activities §2, Activity 2)",
              purpose="Move students from recognizing I-P-O in a given "
                      "example to generating their own.",
              say="With the person next to you, come up with your own "
                  "example — not the calculator or your phone unlocking.",
              expected="A coherent input/process/output triple; cold-call "
                        "2–3 pairs afterward.",
              board="None.",
              transition="Giving a computer instructions to fill that "
                          "Process box is called programming — and it "
                          "starts with a bigger chain than three boxes.")
    return slide


def slide_12_big_question(prs):
    slide = layout_big_question(
        prs, "The Big Question",
        "How do we turn a problem\ninto something a computer can solve?",
        stage_index=3)
    set_notes(slide, markers=["PAUSE", "DO NOT EXPLAIN YET"], timing="1 min",
              purpose="Open the most important idea of the class without "
                      "answering it yet.",
              say="No answer yet — let it sit for a moment.",
              board="None.",
              transition="Everything starts with a real problem. Give me "
                          "one — anything.")
    return slide


def slide_13_pipeline_anchor(prs):
    slide = layout_diagram(prs, "The Core Pipeline", "",
                             draw_single_word("PROBLEM", size=80, color=NAVY),
                             stage_index=3, title_size=1)
    set_notes(slide, markers=["LIVE BOARD", "ASK"],
              timing="13 min (Board Plan Drawing 3, ~43–56 min) — "
              "PROTECTED, never compressed",
              purpose="This slide only anchors the starting word. The "
                      "instructor builds LOGIC → ALGORITHM → PROGRAM "
                      "→ RESULT live on the pen tablet, one box at a "
                      "time, each earned by a question (Board Plan "
                      "Drawing 3, Stages 1–9).",
              question="Before I touch a computer, what's the very first "
                        "thing I have to do?",
              board="LIVE BOARD — the single most important diagram "
                    "of the class. Blank canvas otherwise.",
              transition="(after the full chain is built) This one goes "
                          "in your notes — we'll refer back to it all "
                          "program.")
    return slide


def slide_14_pipeline_reference(prs):
    items = [("PROBLEM", None), ("LOGIC", None), ("ALGORITHM", None),
              ("PROGRAM", None), ("RESULT", None)]
    slide = layout_diagram(prs, "Reference",
                             "Problem → Logic → Algorithm → Program → Result",
                             draw_pipeline(items, height=Inches(1.35),
                                            font_size=16, highlight_last=True),
                             caption="Clean reference, shown after the "
                                     "live build — keep visible for "
                                     "the recap at the end of class.",
                             stage_index=3, title_size=26)
    set_notes(slide, timing="(shown at the end of the live build)",
              purpose="Clean callback of the class's most important "
                      "diagram — also kept visible during the final "
                      "recap (Board Plan §12).",
              board="Matches the finished pen-tablet drawing exactly.",
              transition="This is still just words on a board — how do "
                          "we find the larger number?")
    return slide


def slide_15_build_algorithm(prs):
    slide = layout_activity(
        prs, "Live Problem-Solving", "Build an Algorithm",
        ["A = 15        B = 21", "How would you find the larger number?"],
        sub="Talk me through it — don't just give me the answer.",
        stage_index=4)
    set_notes(slide, markers=["ASK", "LET STUDENTS ANSWER"],
              timing="(nested inside the 13-min pipeline segment)",
              purpose="Students construct the algorithm themselves, live "
                      "(Activities §4, Activity 4).",
              question="Given A=15, B=21, how would you find the larger "
                        "one?",
              expected="'Compare them', 'check which is bigger' — "
                        "press for the method, not just the answer (21).",
              board="LIVE BOARD — steps written under the ALGORITHM "
                    "box of Drawing 3.",
              transition="Let's see the finished version.")
    return slide


def slide_16_algorithm(prs):
    steps = ["Take two numbers.", "Compare them.",
              "If A > B, A is larger.", "Otherwise, B is larger."]
    slide = layout_recap(prs, "The Algorithm",
                           "“Find the Larger of Two Numbers”",
                           steps, numbered=True, stage_index=4)
    set_notes(slide, timing="1 min",
              purpose="Clean final version, shown only after the class "
                      "has built it themselves (Guide §5.6).",
              say="Notice it doesn't mention any programming language at "
                  "all.",
              board="Matches the live-built version.",
              transition="Is an algorithm the same thing as a program?")
    return slide


def slide_17_algorithm_vs_program(prs):
    slide = layout_two_column(
        prs, "The Distinction", "Algorithm ≠ Program",
        {"tag": "THE RECIPE", "title": "Algorithm",
         "sub": "A solution strategy. Language-independent — works "
                "the same in any language, even plain English."},
        {"tag": "THE DISH", "title": "Program",
         "sub": "That strategy implemented in a programming language. "
                "Only a computer executes it directly."},
        stage_index=4)
    set_notes(slide, markers=["ASK"], timing="3 min",
              purpose="Cement the single most important distinction of "
                      "the class (Guide §5.6).",
              question="Could two different languages implement the "
                        "exact same algorithm?",
              expected="Yes — the steps don't change, only the "
                        "language does.",
              board="Two-word note (“recipe/dish”) jotted next to "
                    "Drawing 3, Stage 9 (optional).",
              transition="Before you can turn a problem into an "
                          "algorithm, you need a certain way of thinking.")
    return slide


def slide_18_computational_thinking(prs):
    items = [
        {"title": "Decomposition", "sub": "Break a big problem into "
                                            "smaller pieces."},
        {"title": "Pattern Recognition", "sub": "Find similarities "
                                                  "between problems."},
        {"title": "Abstraction", "sub": "Ignore unnecessary detail."},
        {"title": "Algorithmic Thinking", "sub": "Create clear, ordered "
                                                   "steps."},
    ]
    slide = layout_diagram(
        prs, "How Programmers Think", "How Do Programmers THINK About Problems?",
        draw_quadrants(items), stage_index=5, title_size=28)
    set_notes(slide, markers=["LIVE BOARD", "ASK"],
              timing="8 min (Board Plan Drawing 4, ~59–67 min)",
              purpose="Four thinking habits, filled in one quadrant at a "
                      "time from a question — not a definitions list "
                      "(Guide §5.7).",
              board="LIVE BOARD — instructor fills each quadrant live, "
                    "question first, then word + one-line example.",
              transition="If I asked you to plan a birthday party, would "
                          "you do it as one giant task?")
    return slide


def slide_19_live_thinking_prompt(prs):
    slide = layout_big_question(
        prs, "Live Thinking",
        "Plan a birthday party.\nOne giant task — or smaller pieces?",
        stage_index=5)
    set_notes(slide, markers=["ASK", "LET STUDENTS ANSWER"], timing="2 min",
              purpose="Introduce Decomposition through a relatable "
                      "example (Question Bank Q7.1).",
              expected="“Break it into pieces — invitations, food, "
                        "venue.”",
              board="Feeds directly into Drawing 4's first quadrant.",
              transition="You already used all four of these today "
                          "without realizing it.")
    return slide


def slide_20_tea_activity(prs):
    slide = layout_section_divider(
        prs, "The Main Activity", "Teach a Computer\nHow to Make Tea",
        subtitle="Give me instructions. I will follow EXACTLY what you say.",
        stage_index=6)
    set_notes(slide, markers=["ACTIVITY", "LIVE BOARD"],
              timing="13–15 min — PROTECTED, never compressed",
              purpose="The emotional and pedagogical centerpiece of the "
                      "class. Full run-of-show in "
                      "03_Tea_Activity_Facilitator_Guide.md — this "
                      "slide deliberately shows nothing beyond the "
                      "opening line so nothing is spoiled.",
              say="Imagine I am a computer. I will do exactly — and "
                  "only — what you tell me, in the exact order you "
                  "tell me.",
              board="Switch to camera / minimal screen for the activity "
                    "itself — no slide or board needed.",
              transition="(after the activity) That gap between what you "
                          "meant and what you actually said is exactly "
                          "the lesson.")
    return slide


def slide_21_lesson_from_tea(prs):
    slide = new_slide(prs, bg=NAVY)
    add_kicker(slide, MARGIN, Inches(0.85), CONTENT_W, "The Key Lesson",
                color=CYAN)
    add_text(slide, MARGIN, Inches(1.35), CONTENT_W, Inches(1.1),
              "Computers don't assume.", size=44, color=WHITE, bold=True,
              font=FONT_DISPLAY)
    add_rich_text(slide, MARGIN, Inches(2.75), CONTENT_W, Inches(0.8),
                    [("Human intention  ", WHITE, False),
                     ("≠", CYAN, True),
                     ("  Precise instruction", WHITE, False)],
                    size=30, font=FONT_DISPLAY)
    add_rect(slide, MARGIN, Inches(3.85), Inches(0.85), Inches(0.045),
              fill=CYAN)
    add_text(slide, MARGIN, Inches(4.15), Inches(10.8), Inches(1.3),
              "Programming requires us to turn human intentions\n"
              "into precise instructions.", size=26,
              color=RGBColor(0xE3, 0xEC, 0xFB), font=FONT_BODY,
              line_spacing=1.2)
    add_text(slide, MARGIN, SLIDE_H - Inches(0.42), Inches(6.5), Inches(0.3),
              FOOTER_TEXT, size=9.5, color=RGBColor(0x6B, 0x82, 0xAE))
    add_progress_dots(slide, 6)
    set_notes(slide, timing="2 min",
              purpose="The most memorable line of the class — say it "
                      "once, plainly, don't paraphrase it longer.",
              board="None.",
              transition="That gap between what you meant and what you "
                          "actually said is exactly what separates how "
                          "we've been programming from something newer.")
    return slide


def slide_22_three_ways(prs):
    rows = [
        ("Traditional\nProgramming", ["INPUT", "RULES", "OUTPUT"], None),
        ("Machine\nLearning", ["DATA", "LEARNING", "MODEL", "PREDICTION"], 0),
        ("Generative\nAI", ["PROMPT", "LLM", "RESPONSE"], 0),
    ]
    slide = layout_diagram(prs, "Reference", "Three Ways to Compute",
                             draw_stacked_pipelines(rows),
                             caption="Conceptual only — no architecture, "
                                     "no training details.",
                             stage_index=7, title_size=30)
    set_notes(slide, markers=["LIVE BOARD"],
              timing="4.5 min (Board Plan Drawing 5, ~86–90.5 min)",
              purpose="A light conceptual bridge — not an ML lecture "
                      "(Guide §5.9). Clean reference only; instructor "
                      "builds this live, row by row.",
              say="In traditional programming, a human writes the "
                  "rules. In ML, rules are learned from data. In "
                  "GenAI, a model generates new content from a prompt.",
              board="LIVE BOARD — instructor draws all three rows.",
              transition="Look at the left-most box in each row — "
                          "what changed?")
    return slide


def slide_23_what_changed(prs):
    slide = layout_big_question(
        prs, "Quick Question", "What changed?",
        note="Rules  →  Data  →  Prompt", stage_index=7)
    set_notes(slide, markers=["ASK", "LET STUDENTS ANSWER"], timing="2 min",
              purpose="The core discovery question of Drawing 5 "
                      "(Question Bank Q9.1).",
              question="Look at the left-most box in each row. What "
                        "changed, row by row?",
              expected="It goes from rules, to data, to a prompt.",
              board="Three rows stay visible.",
              transition="If ML and GenAI are newer branches, where do "
                          "they come from? Let's build the whole "
                          "staircase.")
    return slide


def slide_24_staircase(prs):
    labels = ["Programming", "Computer\nScience", "Mathematics", "Data",
               "Machine\nLearning", "Deep\nLearning", "LLMs", "Agents"]
    slide = layout_diagram(prs, "Reference", "The Learning Staircase",
                             draw_staircase(labels, highlight_idx=(0, 1)),
                             caption="A learning pathway, not a strict "
                                     "formal dependency graph.",
                             stage_index=7, title_size=30)
    set_notes(slide, markers=["LIVE BOARD"],
              timing="4.5 min (Board Plan Drawing 6, ~90.5–95 min)",
              purpose="“We build the staircase, we don't jump to the "
                      "top” — physical metaphor built live "
                      "(Guide §5.9).",
              board="LIVE BOARD — instructor draws ascending steps; "
                    "circles steps 1–2 in red.",
              transition="Why don't we start with AI Agents?")
    return slide


def slide_25_why_not_agents(prs):
    slide = new_slide(prs, bg=LIGHT_BG)
    add_kicker(slide, MARGIN, Inches(1.1), CONTENT_W, "Reflect", color=BLUE)
    add_text(slide, Inches(1.1), Inches(1.7), Inches(11.1), Inches(1.3),
              "Why don't we start\nwith AI Agents?", size=38, color=NAVY,
              bold=True, font=FONT_DISPLAY, align=PP_ALIGN.CENTER,
              line_spacing=1.1)
    add_rect(slide, Inches(5.4), Inches(3.55), Inches(2.5), Inches(0.03),
              fill=CYAN)
    add_text(slide, Inches(1.1), Inches(3.9), Inches(11.1), Inches(1.2),
              "Because we're building the staircase,\nnot jumping to "
              "the top.", size=26, color=BLUE, font=FONT_BODY,
              align=PP_ALIGN.CENTER, line_spacing=1.2, italic=True)
    add_footer(slide)
    add_progress_dots(slide, 8)
    set_notes(slide, markers=["ASK", "PAUSE", "REVEAL"], timing="2 min",
              purpose="Motivational but technically grounded close to "
                      "the AI bridge section.",
              question="Why don't we start with AI Agents?",
              board="None.",
              transition="Let's look at the actual four-month journey.")
    return slide


def slide_26_roadmap(prs):
    cards = [
        {"month": "MONTH 1", "theme": "Foundations",
         "items": ["Computing", "Python", "Logic", "Mathematics"]},
        {"month": "MONTH 2", "theme": "CS & Problem Solving",
         "items": ["C", "DSA", "Algorithms", "SQL"]},
        {"month": "MONTH 3", "theme": "Engineering + AI",
         "items": ["Linux", "Git", "APIs", "AI/ML", "GenAI", "Agents"]},
        {"month": "MONTH 4", "theme": "Consolidation",
         "items": ["Revision", "Projects", "Tests", "Demo Day"]},
    ]
    slide = layout_roadmap(prs, "The Path Ahead", "Your 4-Month Journey",
                             cards, stage_index=8)
    set_notes(slide, timing="7 min",
              purpose="Show why the program is sequenced this way, not "
                      "list every topic (Guide §5.10).",
              say="Explain the reasoning: you can't build reliable AI "
                  "systems on shaky programming fundamentals.",
              board="None.",
              transition="Let's check what stuck today.")
    return slide


def slide_27_what_you_should_know(prs):
    items = [
        "Computer Science is broader than coding.",
        "Programs are instructions executed by computers.",
        "Algorithms describe solution strategies.",
        "Algorithms are not the same as programs.",
        "Computers require precise instructions.",
        "Computational thinking helps solve problems.",
        "Programming and CS form foundations for AI.",
    ]
    slide = layout_recap(prs, "Remember This", "What You Should Know Today",
                           items, numbered=False, stage_index=9)
    set_notes(slide, timing="(reference — pairs with the verbal recap)",
              purpose="Seven memorable statements, matching the Student "
                      "Notes 'Remember These' section exactly.",
              board="None.",
              transition="Let's check what stuck today.")
    return slide


def slide_28_exit_check(prs):
    items = [
        "Is Computer Science the same as coding?",
        "Give an input → process → output example.",
        "Describe an algorithm.",
        "Algorithm vs. program — what's the difference?",
        "How does today's class connect to AI?",
    ]
    slide = layout_recap(prs, "Final Recap", "Exit Check", items,
                           numbered=True, stage_index=9)
    set_notes(slide, markers=["ASK", "LET STUDENTS ANSWER"],
              timing="6–8 min (Guide §11, Exit Questions §06)",
              purpose="Test understanding, not memorization — accept "
                      "answers in the student's own words.",
              board="Drawing 3 and the Staircase (Drawing 6) stay "
                    "visible for reference (Board Plan §12).",
              transition="One last thing before you go — your first "
                          "assignment.")
    return slide


def slide_29_homework(prs):
    activities = ["ATM withdrawal", "Ordering food", "Booking a cab",
                   "Online shopping", "Making tea", "Sending a message"]
    steps = ["Identify Input, Process, and Output.",
              "Write an algorithm with at least 5 steps.",
              "Use plain language — no code required."]
    slide = layout_homework(prs, "Assignment 1", "Think Like a Programmer",
                              "Choose one everyday process from the list "
                              "below.", activities, steps, stage_index=9)
    set_notes(slide, timing="2 min",
              purpose="Close the class with a clear, achievable task "
                      "(Student Notes §Homework, Interaction Pack "
                      "§06).",
              say="Full instructions, format, and a challenge version are "
                  "in your Student Notes handout.",
              board="None.",
              transition="Class ends.")
    return slide


# ============================================================================
# 6. MAIN GENERATION FUNCTION
# ============================================================================

SLIDE_BUILDERS = [
    slide_01_title, slide_02_opening_question, slide_03_what_is_cs,
    slide_04_cs_not_coding, slide_05_cs_map, slide_06_quick_question,
    slide_07_button_press, slide_08_three_words, slide_09_quick_check,
    slide_10_ipo, slide_11_your_turn, slide_12_big_question,
    slide_13_pipeline_anchor, slide_14_pipeline_reference,
    slide_15_build_algorithm, slide_16_algorithm,
    slide_17_algorithm_vs_program, slide_18_computational_thinking,
    slide_19_live_thinking_prompt, slide_20_tea_activity,
    slide_21_lesson_from_tea, slide_22_three_ways, slide_23_what_changed,
    slide_24_staircase, slide_25_why_not_agents, slide_26_roadmap,
    slide_27_what_you_should_know, slide_28_exit_check, slide_29_homework,
]


def build_presentation(output_path):
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    for i, builder in enumerate(SLIDE_BUILDERS, start=1):
        slide = builder(prs)
        # page number on every slide except the title slide
        if i != 1:
            add_page_number(slide, i)
    prs.save(output_path)
    return len(SLIDE_BUILDERS)


if __name__ == "__main__":
    here = os.path.dirname(os.path.abspath(__file__))
    out = os.path.join(here, "Class_01_Welcome_to_Computer_Science.pptx")
    n = build_presentation(out)
    print(f"Generated {n} slides -> {out}")
