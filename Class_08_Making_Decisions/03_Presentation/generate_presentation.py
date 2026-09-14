# -*- coding: utf-8 -*-
"""
generate_presentation.py

Generates Class_08_Making_Decisions.pptx for
Bong Study Hub - Foundation Batch 2026.

Source of truth (do not redesign the lesson when editing this script):
  01_Master_Instructor_Guide/Class_08_Master_Instructor_Guide.md
  02_Student_Notes/Class_08_Student_Notes.md

This deck is a VISUAL TEACHING AID for a live instructor, not the
Student Notes placed onto slides. Slides carry diagrams, short
statements, questions, and live code the instructor points to and talks
around - never paragraphs read aloud.

Design system, component library, and file structure follow the
convention established in Classes 01-07's generate_presentation.py (same
palette, fonts, canvas, and shape-drawing primitives). Like Class 07,
this class needs **no new diagram shapes** - the side-by-side
code+output comparison built up across Classes 06-07 is exactly the
right shape for "run the same program twice with different data" (the
if/else True-vs-False contrast, and the correctly-indented-vs-
IndentationError contrast), which is this class's signature move.

Structure of this file:
  1. Theme configuration (colors, fonts, spacing)
  2. Low-level utility functions (shapes, text, arrows)
  3. Reusable slide-layout components (title, question, diagram, ...)
  4. Content-aware diagram builders (pipeline, vertical flow, code, ...)
  5. Slide-specific builders (slide_01 .. slide_19)
  6. Main generation function
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR

# ============================================================================
# 1. THEME CONFIGURATION
# ============================================================================

# -- Colors ---------------------------------------------------------------
# Same palette as Classes 01-07, for cross-class brand consistency.
NAVY = RGBColor(0x1E, 0x3A, 0x72)
BLUE = RGBColor(0x2A, 0x5C, 0xB8)
CYAN = RGBColor(0x22, 0xC7, 0xE0)
INK = RGBColor(0x1F, 0x27, 0x33)
MUTED = RGBColor(0x64, 0x70, 0x85)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_BG = RGBColor(0xF7, 0xF9, 0xFC)
CARD_BG = RGBColor(0xEE, 0xF3, 0xFB)
CARD_BORDER = RGBColor(0xC7, 0xD6, 0xEC)
CYAN_BG = RGBColor(0xE4, 0xF9, 0xFB)
CYAN_BORDER = RGBColor(0x9F, 0xE6, 0xEE)

# Carried over from Classes 05-07: an error/warning accent, used this
# class for a real, deliberate IndentationError (Section 12).
ERROR_RED = RGBColor(0xB8, 0x3A, 0x3A)
ERROR_BG = RGBColor(0xFC, 0xEC, 0xEC)
ERROR_BORDER = RGBColor(0xEE, 0xB9, 0xB9)

# -- Fonts ------------------------------------------------------------------
FONT_DISPLAY = "Aptos Display"
FONT_BODY = "Aptos"
# Carried over from Class 05: a monospace face for literal Python source
# code or program output.
FONT_MONO = "Consolas"

# -- Canvas -----------------------------------------------------------------
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
MARGIN = Inches(0.7)
CONTENT_W = SLIDE_W - 2 * MARGIN

FOOTER_TEXT = "Bong Study Hub  ·  Foundation Batch 2026  ·  Class 08"


# ============================================================================
# 2. LOW-LEVEL UTILITIES
# ============================================================================

def new_slide(prs, bg=LIGHT_BG):
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = bg
    return slide


def _i(v):
    return int(round(v))


def _no_shadow(shape):
    try:
        shape.shadow.inherit = False
    except Exception:
        pass


def add_rect(slide, x, y, w, h, fill=WHITE, line=None, line_w=Pt(1),
             radius=0.06, shadow=False):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shp = slide.shapes.add_shape(shape_type, _i(x), _i(y), _i(w), _i(h))
    if radius:
        try:
            shp.adjustments[0] = radius
        except Exception:
            pass
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = line_w
    if not shadow:
        _no_shadow(shp)
    shp.text_frame.margin_left = 0
    shp.text_frame.margin_right = 0
    shp.text_frame.margin_top = 0
    shp.text_frame.margin_bottom = 0
    return shp


def add_oval(slide, x, y, w, h, fill=WHITE, line=None, line_w=Pt(1)):
    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, _i(x), _i(y), _i(w), _i(h))
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = line_w
    _no_shadow(shp)
    return shp


def add_text(slide, x, y, w, h, text, size=18, color=INK, bold=False,
             italic=False, align=PP_ALIGN.LEFT, font=FONT_BODY,
             anchor=MSO_ANCHOR.TOP, line_spacing=1.08, wrap=True,
             shrink=False):
    box = slide.shapes.add_textbox(_i(x), _i(y), _i(w), _i(h))
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    if shrink:
        tf.auto_size = MSO_AUTO_SIZE.NONE
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.name = font
        run.font.color.rgb = color
    return box


def add_line(slide, x1, y1, x2, y2, color=BLUE, width=Pt(1.5)):
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, _i(x1), _i(y1),
                                        _i(x2), _i(y2))
    conn.line.color.rgb = color
    conn.line.width = width
    _no_shadow(conn)
    return conn


def add_triangle(slide, cx, cy, size, color, rotation):
    s = _i(size)
    shp = slide.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE,
                                  _i(cx) - s // 2, _i(cy) - s // 2, s, s)
    shp.rotation = rotation
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.line.fill.background()
    _no_shadow(shp)
    return shp


def add_arrow_h(slide, x1, y, x2, color=BLUE, width=Pt(1.5)):
    tip_gap = Emu(90000)
    add_line(slide, x1, y, x2 - tip_gap, y, color=color, width=width)
    add_triangle(slide, x2 - Emu(45000), y, Emu(160000), color, rotation=90)


def add_arrow_v(slide, x, y1, y2, color=BLUE, width=Pt(1.5)):
    tip_gap = Emu(90000)
    add_line(slide, x, y1, x, y2 - tip_gap, color=color, width=width)
    add_triangle(slide, x, y2 - Emu(45000), Emu(160000), color, rotation=180)


def add_footer(slide):
    add_text(slide, MARGIN, SLIDE_H - Inches(0.42), Inches(7.5), Inches(0.3),
              FOOTER_TEXT, size=9.5, color=MUTED, font=FONT_BODY)


def add_page_number(slide, number):
    add_text(slide, SLIDE_W - Inches(1.1), SLIDE_H - Inches(0.42),
              Inches(0.5), Inches(0.3), str(number), size=9.5, color=MUTED,
              align=PP_ALIGN.RIGHT, font=FONT_BODY)


def add_kicker(slide, x, y, w, text, color=BLUE):
    add_text(slide, x, y, w, Inches(0.4), text.upper(), size=13, color=color,
              bold=True, font=FONT_BODY)


def set_notes(slide, purpose=None, say=None, question=None, expected=None,
               board=None, transition=None, timing=None, markers=None):
    lines = []
    if markers:
        lines.append("[" + "] [".join(markers) + "]")
    if timing:
        lines.append(f"TIME: {timing}")
    if purpose:
        lines.append(f"Purpose: {purpose}")
    if say:
        lines.append(f"Say: {say}")
    if question:
        lines.append(f"Ask: {question}")
    if expected:
        lines.append(f"Expected: {expected}")
    if board:
        lines.append(f"Board: {board}")
    if transition:
        lines.append(f"Transition: {transition}")
    slide.notes_slide.notes_text_frame.text = "\n".join(lines)


# ============================================================================
# 3. REUSABLE SLIDE-LAYOUT COMPONENTS
# ============================================================================

def layout_title(prs, kicker, title_lines, subtitle):
    slide = new_slide(prs, bg=NAVY)
    add_rect(slide, MARGIN, Inches(2.15), Inches(1.2), Inches(0.06), fill=CYAN)
    add_kicker(slide, MARGIN, Inches(1.75), Inches(10), kicker, color=CYAN)
    y = Inches(2.4)
    for line in title_lines:
        add_text(slide, MARGIN, y, Inches(11.5), Inches(1.1), line, size=50,
                  color=WHITE, bold=True, font=FONT_DISPLAY)
        y += Inches(1.0)
    add_text(slide, MARGIN, y + Inches(0.15), Inches(10.5), Inches(0.6),
              subtitle, size=20, color=RGBColor(0xB9, 0xCC, 0xEE),
              italic=True, font=FONT_BODY)
    add_text(slide, MARGIN, SLIDE_H - Inches(0.9), Inches(6), Inches(0.5),
              "Bong Study Hub\nFoundation Batch 2026", size=13,
              color=RGBColor(0x8F, 0xA6, 0xD1), font=FONT_BODY,
              line_spacing=1.15)
    return slide


def layout_section_divider(prs, kicker, title, subtitle=None, title_size=40):
    slide = new_slide(prs, bg=NAVY)
    add_rect(slide, MARGIN, Inches(3.05), Inches(1.0), Inches(0.06), fill=CYAN)
    if kicker:
        add_kicker(slide, MARGIN, Inches(2.65), Inches(10), kicker, color=CYAN)
    add_text(slide, MARGIN, Inches(3.25), Inches(11.8), Inches(2.4), title,
              size=title_size, color=WHITE, bold=True, font=FONT_DISPLAY,
              line_spacing=1.15)
    if subtitle:
        add_text(slide, MARGIN, Inches(5.6), Inches(10.5), Inches(0.8),
                  subtitle, size=18, color=RGBColor(0xB9, 0xCC, 0xEE),
                  font=FONT_BODY, italic=True)
    add_text(slide, MARGIN, SLIDE_H - Inches(0.42), Inches(7.5), Inches(0.3),
              FOOTER_TEXT, size=9.5, color=RGBColor(0x6B, 0x82, 0xAE))
    return slide


def _content_header(slide, kicker, title, title_size=32):
    if kicker:
        add_kicker(slide, MARGIN, Inches(0.5), CONTENT_W, kicker, color=BLUE)
        title_y = Inches(0.86)
    else:
        title_y = Inches(0.6)
    add_text(slide, MARGIN, title_y, CONTENT_W, Inches(0.9), title,
              size=title_size, color=NAVY, bold=True, font=FONT_DISPLAY)
    add_rect(slide, MARGIN, title_y + Inches(0.85), Inches(0.85),
              Inches(0.045), fill=CYAN)
    return title_y + Inches(1.15)


def layout_big_question(prs, kicker, question, note=None, size=38):
    slide = new_slide(prs, bg=LIGHT_BG)
    if kicker:
        add_kicker(slide, MARGIN, Inches(1.5), CONTENT_W, kicker, color=BLUE)
    add_text(slide, Inches(1.0), Inches(2.3), Inches(11.3), Inches(2.9),
              question, size=size, color=NAVY, bold=True, font=FONT_DISPLAY,
              align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
              line_spacing=1.15)
    if note:
        add_text(slide, Inches(1.0), Inches(5.5), Inches(11.3), Inches(0.5),
                  note, size=15, color=MUTED, align=PP_ALIGN.CENTER,
                  italic=True)
    add_footer(slide)
    return slide


def layout_diagram(prs, kicker, title, diagram_fn, caption=None,
                    title_size=32):
    slide = new_slide(prs, bg=LIGHT_BG)
    y = _content_header(slide, kicker, title, title_size=title_size)
    area_y = y + Inches(0.25)
    area_h = SLIDE_H - area_y - Inches(0.75)
    if caption:
        area_h -= Inches(0.5)
    diagram_fn(slide, MARGIN, area_y, CONTENT_W, area_h)
    if caption:
        add_text(slide, MARGIN, SLIDE_H - Inches(1.05), CONTENT_W,
                  Inches(0.4), caption, size=13, color=MUTED, italic=True,
                  align=PP_ALIGN.CENTER)
    add_footer(slide)
    return slide


def layout_full_screen_visual(prs, kicker, title, diagram_fn, caption=None):
    slide = new_slide(prs, bg=LIGHT_BG)
    add_kicker(slide, MARGIN, Inches(0.45), CONTENT_W, kicker, color=BLUE)
    if title:
        add_text(slide, MARGIN, Inches(0.8), CONTENT_W, Inches(0.7), title,
                  size=26, color=NAVY, bold=True, font=FONT_DISPLAY)
    diagram_fn(slide, MARGIN, Inches(1.55), CONTENT_W,
                SLIDE_H - Inches(1.55) - Inches(1.15))
    if caption:
        add_text(slide, MARGIN, SLIDE_H - Inches(1.0), CONTENT_W, Inches(0.4),
                  caption, size=13, color=MUTED, italic=True,
                  align=PP_ALIGN.CENTER)
    add_footer(slide)
    return slide


def layout_activity(prs, kicker, title, prompt_lines, sub=None, tag="ACTIVITY"):
    slide = new_slide(prs, bg=LIGHT_BG)
    add_rect(slide, Inches(0.9), Inches(0.55), Inches(1.5), Inches(0.42),
              fill=CYAN_BG, line=CYAN_BORDER, line_w=Pt(1))
    add_text(slide, Inches(0.9), Inches(0.55), Inches(1.5), Inches(0.42),
              tag, size=13, color=NAVY, bold=True, align=PP_ALIGN.CENTER,
              anchor=MSO_ANCHOR.MIDDLE)
    add_kicker(slide, Inches(2.6), Inches(0.62), Inches(9), kicker,
                color=BLUE)
    add_text(slide, Inches(0.9), Inches(1.25), Inches(11.5), Inches(1.0),
              title, size=30, color=NAVY, bold=True, font=FONT_DISPLAY)
    y = Inches(2.5)
    for line in prompt_lines:
        add_text(slide, Inches(0.9), y, Inches(11.5), Inches(0.85), line,
                  size=24, color=INK, font=FONT_BODY, line_spacing=1.15)
        y += Inches(0.8)
    if sub:
        add_text(slide, Inches(0.9), y + Inches(0.15), Inches(11.0),
                  Inches(0.6), sub, size=15, color=MUTED, italic=True)
    add_footer(slide)
    return slide


# ============================================================================
# 4. CONTENT-AWARE DIAGRAM BUILDERS
# ============================================================================

def draw_vertical_flow(items, highlight_indices=None, box_w=Inches(4.6),
                        font_size=16):
    highlight_indices = highlight_indices or set()

    def _fn(slide, x, y, w, h):
        n = len(items)
        box_h = min(Inches(0.75), h / (n * 1.35))
        gap = (h - n * box_h) / (n - 1) if n > 1 else Inches(0)
        gap = max(Inches(0.1), min(gap, Inches(0.5)))
        total_h = n * box_h + (n - 1) * gap
        cy = y + max(Inches(0), (h - total_h) / 2)
        bx = x + w / 2 - box_w / 2
        for i, label in enumerate(items):
            is_hl = i in highlight_indices
            fill = CYAN_BG if is_hl else CARD_BG
            border = CYAN_BORDER if is_hl else CARD_BORDER
            add_rect(slide, bx, cy, box_w, box_h, fill=fill, line=border,
                      line_w=Pt(1.4 if is_hl else 1.1), radius=0.14)
            add_text(slide, bx, cy, box_w, box_h, label, size=font_size,
                      color=NAVY, bold=True, align=PP_ALIGN.CENTER,
                      anchor=MSO_ANCHOR.MIDDLE)
            if i < n - 1:
                add_arrow_v(slide, x + w / 2, cy + box_h + Inches(0.03),
                            cy + box_h + gap - Inches(0.03), color=BLUE,
                            width=Pt(1.75))
            cy += box_h + gap
    return _fn


def draw_pipeline(items, height=Inches(1.3), font_size=15,
                   highlight_indices=None, gap=Inches(0.55)):
    highlight_indices = highlight_indices or set()

    def _fn(slide, x, y, w, h):
        n = len(items)
        box_w = (w - gap * (n - 1)) / n
        box_y = y + (h - height) / 2
        cx = x
        for i, (label, sub) in enumerate(items):
            is_hl = i in highlight_indices
            fill = CYAN_BG if is_hl else CARD_BG
            border = CYAN_BORDER if is_hl else CARD_BORDER
            add_rect(slide, cx, box_y, box_w, height, fill=fill,
                      line=border, line_w=Pt(1.25), radius=0.09)
            label_y = box_y + (height / 2 - Inches(0.32) if sub
                                else height / 2 - Inches(0.24))
            add_text(slide, cx, label_y, box_w, Inches(0.5), label,
                      size=font_size, color=NAVY, bold=True,
                      align=PP_ALIGN.CENTER, font=FONT_BODY)
            if sub:
                add_text(slide, cx, label_y + Inches(0.42), box_w,
                          Inches(0.35), sub, size=font_size - 4, color=MUTED,
                          align=PP_ALIGN.CENTER)
            if i < n - 1:
                add_arrow_h(slide, cx + box_w + Inches(0.06), box_y + height / 2,
                            cx + box_w + gap - Inches(0.06), color=BLUE,
                            width=Pt(1.75))
            cx += box_w + gap
    return _fn


def draw_word_grid(words, cols=4, highlight_indices=None, font_size=20):
    highlight_indices = highlight_indices or set()

    def _fn(slide, x, y, w, h):
        n = len(words)
        rows = (n + cols - 1) // cols
        gap = Inches(0.3)
        cell_w = (w - gap * (cols - 1)) / cols
        cell_h = min(Inches(1.0), (h - gap * (rows - 1)) / rows) if rows else h
        grid_h = rows * cell_h + (rows - 1) * gap
        top = y + max(Inches(0), (h - grid_h) / 2)
        for i, word in enumerate(words):
            r, c = divmod(i, cols)
            last_row = (r == rows - 1)
            items_in_row = cols if (not last_row or n % cols == 0) else n % cols
            row_w = items_in_row * cell_w + (items_in_row - 1) * gap
            row_x0 = x + (w - row_w) / 2
            cx = row_x0 + c * (cell_w + gap)
            cy = top + r * (cell_h + gap)
            is_hl = i in highlight_indices
            fill = CYAN_BG if is_hl else CARD_BG
            border = CYAN_BORDER if is_hl else CARD_BORDER
            add_rect(slide, cx, cy, cell_w, cell_h, fill=fill, line=border,
                      line_w=Pt(1.25), radius=0.14)
            add_text(slide, cx, cy, cell_w, cell_h, word, size=font_size,
                      color=NAVY, bold=True, align=PP_ALIGN.CENTER,
                      anchor=MSO_ANCHOR.MIDDLE, font=FONT_BODY)
    return _fn


def draw_definition(term, definition, term_size=72, def_size=28):
    def _fn(slide, x, y, w, h):
        add_text(slide, x, y, w, h * 0.42, term, size=term_size, color=NAVY,
                  bold=True, font=FONT_DISPLAY, align=PP_ALIGN.CENTER,
                  anchor=MSO_ANCHOR.MIDDLE)
        rule_y = y + h * 0.46
        add_rect(slide, x + w / 2 - Inches(0.5), rule_y, Inches(1.0),
                  Inches(0.04), fill=CYAN)
        add_text(slide, x, rule_y + Inches(0.25), w, h * 0.5 - Inches(0.25),
                  definition, size=def_size, color=INK, bold=True,
                  align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.TOP,
                  line_spacing=1.2, font=FONT_BODY)
    return _fn


# -- Carried over from Classes 05-07: code / output shapes ------------------

# New for Class 08: a lighter navy tint used only for the indent-guide
# rule in draw_code_block - subtle against the dark editor card, but a
# real visual signal in a class that's specifically about indentation.
INDENT_GUIDE = RGBColor(0x3B, 0x57, 0x8E)


def draw_code_block(lines, font_size=22, indent_unit=Inches(0.36)):
    """A dark 'editor' card holding monospace Python source, one string per
    line. The small cyan dot in the corner reads as a window/tab marker.

    Leading spaces on each line (Python's actual block-nesting syntax)
    are read off and converted into an explicit indent level: the text
    itself is shifted right by that many indent_units, and a thin
    vertical guide line is drawn at each indent stop - an IDE-style
    indent guide. This keeps indentation visually unambiguous even
    where a viewer's rendering substitutes a non-monospace font (leading
    space *characters* alone can render too narrow to read as an
    indent) - important specifically because Class 08 is about
    indentation as syntax, not just about the code's content."""
    def _fn(slide, x, y, w, h):
        n = max(1, len(lines))
        card_h = min(h, Inches(0.62) * n + Inches(0.55))
        top = y + max(Inches(0), (h - card_h) / 2)
        add_rect(slide, x, top, w, card_h, fill=NAVY, radius=0.06)
        add_oval(slide, x + Inches(0.32), top + Inches(0.28), Inches(0.14),
                  Inches(0.14), fill=CYAN, line=None)
        line_h = (card_h - Inches(0.55)) / n
        ty = top + Inches(0.48)
        base_x = x + Inches(0.62)

        parsed = []
        max_level = 0
        for line in lines:
            stripped = line.lstrip(" ")
            level = (len(line) - len(stripped)) // 4
            parsed.append((level, stripped))
            max_level = max(max_level, level)

        for lvl in range(1, max_level + 1):
            gx = base_x + (lvl - 1) * indent_unit + Inches(0.09)
            add_line(slide, gx, top + Inches(0.12), gx,
                      top + card_h - Inches(0.12), color=INDENT_GUIDE,
                      width=Pt(1.25))

        for i, (level, text) in enumerate(parsed):
            tx = base_x + level * indent_unit
            add_text(slide, tx, ty + i * line_h, w - Inches(0.95) - level * indent_unit,
                      line_h, text, size=font_size, color=WHITE,
                      font=FONT_MONO, anchor=MSO_ANCHOR.MIDDLE)
    return _fn


def draw_output_card(lines, is_error=False, font_size=20, label=None):
    """A light 'console' card for program output, or an error-tinted card
    when is_error is True."""
    def _fn(slide, x, y, w, h):
        fill = ERROR_BG if is_error else CARD_BG
        border = ERROR_BORDER if is_error else CARD_BORDER
        bar = ERROR_RED if is_error else CYAN
        text_color = ERROR_RED if is_error else INK
        add_rect(slide, x, y, w, h, fill=fill, line=border, line_w=Pt(1.25),
                  radius=0.05)
        add_rect(slide, x, y, Inches(0.09), h, fill=bar, radius=None)
        lab = label or ("ERROR" if is_error else "OUTPUT")
        add_text(slide, x + Inches(0.35), y + Inches(0.16),
                  w - Inches(0.6), Inches(0.32), lab, size=12.5,
                  color=(ERROR_RED if is_error else MUTED), bold=True)
        n = max(1, len(lines))
        body_y = y + Inches(0.55)
        body_h = h - Inches(0.7)
        add_text(slide, x + Inches(0.35), body_y, w - Inches(0.6), body_h,
                  "\n".join(lines), size=font_size, color=text_color,
                  font=FONT_MONO, anchor=MSO_ANCHOR.TOP, line_spacing=1.3)
    return _fn


def draw_code_and_output(code_lines, output_lines, is_error=False,
                          code_font=22, output_font=18, output_label=None):
    """Editor card on top, console result (or error) beneath it."""
    def _fn(slide, x, y, w, h):
        gap = Inches(0.35)
        code_h = h * 0.54
        out_h = h - code_h - gap
        draw_code_block(code_lines, font_size=code_font)(slide, x, y, w, code_h)
        draw_output_card(output_lines, is_error=is_error,
                          font_size=output_font, label=output_label)(
            slide, x, y + code_h + gap, w, out_h)
    return _fn


def draw_two_code_examples(left, right, font_size=15, output_font=14):
    """Two side-by-side code+output panels - this class's signature shape:
    the exact same program, run twice with different data, side by side."""
    def _fn(slide, x, y, w, h):
        gap = Inches(0.55)
        col_w = (w - gap) / 2
        for i, spec in enumerate((left, right)):
            cx = x + i * (col_w + gap)
            draw_code_and_output(
                spec["code"], spec["output"], is_error=spec.get("is_error", False),
                code_font=spec.get("code_font", font_size),
                output_font=spec.get("output_font", output_font),
                output_label=spec.get("label"))(slide, cx, y, col_w, h)
    return _fn


# ============================================================================
# 5. SLIDE-SPECIFIC BUILDERS
# ============================================================================

def slide_01_title(prs):
    slide = layout_title(
        prs, "Bong Study Hub · Foundation Batch 2026 · Class 08",
        ["Making Decisions:", "if / elif / else"],
        "The First Code That “Thinks” Before Acting")
    set_notes(slide, timing="2 min",
              purpose="Set tone; do not read the slide aloud.",
              say="Open with the live 'this line prints no matter what' "
                  "demonstration (Guide §8) before this slide is even "
                  "shown, if possible.",
              transition="How do we make a program actually behave "
                          "differently depending on what's true?")
    return slide


def slide_02_big_question(prs):
    slide = layout_big_question(
        prs, "The Big Question",
        "How do we make a program\nactually behave differently\n"
        "depending on what's true?", size=34)
    set_notes(slide, markers=["ASK", "PAUSE", "DO NOT ANSWER YET"],
              timing="(closes Guide §8-9)",
              purpose="Land the central question after the opening "
                      "demo and Class 07 callback.",
              board="None.",
              transition="Let's name the fix.")
    return slide


def slide_03_what_is_a_conditional(prs):
    slide = layout_diagram(
        prs, "Naming the Fix", "What Is a Conditional?",
        draw_definition("CONDITIONAL", "RUNS A BLOCK OF CODE\nONLY IF A "
                          "CONDITION IS TRUE", term_size=54, def_size=26))
    set_notes(slide, timing="5 min (Guide §10)",
              purpose="Name the fix before showing syntax - a "
                      "conditional runs a block only if its condition "
                      "is True.",
              question="What is a conditional, in your own words?",
              expected="Code that only runs a block if a condition is "
                        "true.",
              board="None.",
              transition="Let's write one.")
    return slide


def slide_04_first_if(prs):
    slide = layout_diagram(
        prs, "Same Code, Two Behaviors", "Your First if Statement",
        draw_two_code_examples(
            {"code": ["is_raining = True", "", "if is_raining:",
                       '    print("Bring an umbrella!")'],
             "output": ["Bring an umbrella!"], "label": "is_raining = True"},
            {"code": ["is_raining = False", "", "if is_raining:",
                       '    print("Bring an umbrella!")'],
             "output": ["(nothing prints)"], "label": "is_raining = False"},
            font_size=15))
    set_notes(slide, markers=["TYPE-ALONG"], timing="9 min (Guide §11) "
              "— NON-NEGOTIABLE",
              purpose="Type this live, run it once as True, then change "
                      "the value and run again - the single most "
                      "important moment of the class.",
              say="The indented line belongs inside the if. Same code, "
                  "different data, completely different behavior — "
                  "that's never happened before in this course.",
              question="Before I run it the second time — what do you "
                        "think happens?",
              expected="Nothing prints — the condition is False.",
              board="Live screen.",
              transition="Look at that indented line very closely.")
    return slide


def slide_05_indentation(prs):
    slide = layout_diagram(
        prs, "Break It On Purpose", "Indentation: Not Just Style",
        draw_two_code_examples(
            {"code": ["if is_raining:", '    print("Bring an umbrella!")'],
             "output": ["Bring an umbrella!"], "label": "CORRECTLY INDENTED"},
            {"code": ["if is_raining:", 'print("Bring an umbrella!")'],
             "output": ["IndentationError:", "expected an indented block"],
             "is_error": True, "label": "MISSING INDENT"},
            font_size=15.5))
    set_notes(slide, markers=["LIVE", "BREAK ON PURPOSE"], timing="8 min "
              "(Guide §12) — NON-NEGOTIABLE",
              purpose="Python uses indentation itself as the syntax "
                      "that marks a block - not curly braces, not "
                      "'end'. Break it on purpose, read the error "
                      "calmly.",
              say="Every line in the same block needs the SAME "
                  "indentation. Let your editor auto-indent after a "
                  "colon, and stay consistent.",
              question="What does indentation actually mean in Python?",
              expected="It's the syntax that shows which lines belong "
                        "to a block — not decoration.",
              board="Live screen.",
              transition="What if we also want something to happen "
                          "when the condition is false?")
    return slide


def slide_06_else(prs):
    slide = layout_diagram(
        prs, "Exactly One of Two Paths", "else: The Otherwise",
        draw_two_code_examples(
            {"code": ["is_raining = False", "if is_raining:",
                       '    print("Umbrella!")', "else:",
                       '    print("Sunshine!")'],
             "output": ["Sunshine!"], "label": "is_raining = False"},
            {"code": ["is_raining = True", "if is_raining:",
                       '    print("Umbrella!")', "else:",
                       '    print("Sunshine!")'],
             "output": ["Umbrella!"], "label": "is_raining = True"},
            font_size=14))
    set_notes(slide, markers=["TYPE-ALONG"], timing="7 min (Guide §13) "
              "— NON-NEGOTIABLE",
              purpose="else means 'otherwise' - its block runs only "
                      "when the if condition was False. It never takes "
                      "its own condition.",
              question="When does the else block run? Does else ever "
                        "have its own condition?",
              expected="Only when the if condition was False. No — "
                        "never.",
              board="Live screen.",
              transition="What if there are more than two possible "
                          "paths?")
    return slide


def slide_07_elif(prs):
    slide = layout_diagram(
        prs, "Building Live", "elif: More Than Two Paths",
        draw_code_block(
            ["age = 15", "", "if age < 13:", '    print("Child ticket")',
             "elif age < 20:", '    print("Teen ticket")',
             "elif age < 65:", '    print("Adult ticket")', "else:",
             '    print("Senior ticket")'], font_size=17),
        caption='With age = 15 → "Teen ticket" prints. Change age and '
                're-run to see the other branches.')
    set_notes(slide, markers=["LIVE"], timing="9 min (Guide §14) — "
              "NON-NEGOTIABLE",
              purpose="elif = 'else if' - checked only if every "
                      "condition above it was False. Re-run with "
                      "several different ages, predicting each time.",
              question="What does elif stand for? With age = 15, which "
                        "branch runs?",
              expected="Else if. \"Teen ticket.\"",
              board="Live screen.",
              transition="Let's see exactly what happens when more "
                          "than one condition would technically be "
                          "true.")
    return slide


def slide_08_order_question(prs):
    slide = layout_big_question(
        prs, "Predict Before You Reveal",
        "score is 95 — that's true for BOTH\nconditions below. Which "
        "message\nprints?", size=28,
        note='if score >= 60: "Pass"   ·   elif score >= 90: "Pass with '
             'Distinction"')
    set_notes(slide, markers=["ASK", "PAUSE"], timing="(opens Guide "
              "§15)",
              purpose="Let guesses happen before revealing - many will "
                      "guess 'Pass with Distinction,' since it feels "
                      "more specific/correct.",
              board="None.",
              transition="Let's actually run it and see.")
    return slide


def slide_09_order_reveal(prs):
    slide = layout_diagram(
        prs, "Order Matters", "Only the First True Branch Runs",
        draw_code_and_output(
            ["score = 95", "", "if score >= 60:", '    print("Pass")',
             "elif score >= 90:",
             '    print("Pass with Distinction")'],
            ["Pass"], code_font=17, output_font=22,
            output_label="OUTPUT — the elif below is never even checked"))
    set_notes(slide, timing="(closes Guide §15)",
              purpose="Python checked the first condition, found it "
                      "True, ran that block, and never even looked at "
                      "the elif below - even though it was also true.",
              say="Order matters. The more specific condition needed "
                  "to come first.",
              board="Live screen.",
              transition="Let's bring back last class's toolkit.")
    return slide


def slide_10_comparisons_logic(prs):
    slide = layout_diagram(
        prs, "Reusing Class 07", "Comparisons and Logic Inside a Condition",
        draw_code_and_output(
            ["age = 16", "is_student = True", "",
             "if age >= 13 and age <= 19 and is_student:",
             '    print("Eligible for the student teen discount!")'],
            ["Eligible for the student teen discount!"], code_font=15.5,
            output_font=17))
    set_notes(slide, timing="8 min (Guide §16)",
              purpose="An if doesn't need anything new here - it just "
                      "needs a boolean expression, however it's built.",
              question="What kind of value does an if's condition need "
                        "to be?",
              expected="A boolean — True or False.",
              board="Live screen.",
              transition="Let's step back and look at the three shapes "
                          "we've built today.")
    return slide


def slide_11_common_patterns(prs):
    slide = layout_diagram(
        prs, "Three Shapes", "Common Structure Patterns",
        draw_pipeline(
            [("IF ALONE", "do something, or nothing"),
             ("IF / ELSE", "exactly one of two paths"),
             ("IF / ELIF / ELSE", "exactly one of several paths")],
            height=Inches(1.6), font_size=19, gap=Inches(0.5)))
    set_notes(slide, timing="6 min (Guide §17)",
              purpose="Every conditional written today is one of these "
                      "three shapes - choosing between them is about "
                      "how many outcomes the problem actually has.",
              question="If a problem has exactly two outcomes, which "
                        "shape fits? Three or more?",
              expected="if/else. if/elif/.../else.",
              board="None.",
              transition="Let's connect this whole idea to something "
                          "from Class 01.")
    return slide


def slide_12_class01_callback(prs):
    slide = layout_diagram(
        prs, "Reconnecting to Class 01", "Algorithms Can Branch",
        draw_vertical_flow(
            ["ALGORITHM  (Class 01 — imagined as a straight line)",
             "ALGORITHMS CAN BRANCH  (today — different paths for "
             "different data)"],
            highlight_indices={1}, box_w=Inches(9.6), font_size=16))
    set_notes(slide, timing="7 min (Guide §18)",
              purpose="Class 01's ALGORITHM was always a straight line "
                      "of steps. Today that picture gets wider.",
              question="How did Class 01 describe an algorithm? What "
                        "does today add to that picture?",
              expected="Steps a person could follow, in order. "
                        "Algorithms can branch, not just proceed "
                        "straight through.",
              board="None.",
              transition="Let's put the entire day's reasoning on one "
                          "board — starting all the way back at Class "
                          "01.")
    return slide


def slide_13_full_bridge(prs):
    items = ["PROBLEM", "LOGIC", "ALGORITHM", "PYTHON CODE", "VARIABLES",
             "EXPRESSIONS", "DECISIONS", "RUN", "OUTPUT / RESULT"]
    slide = layout_full_screen_visual(
        prs, "The Hero Diagram", None,
        draw_vertical_flow(items, highlight_indices={6}, box_w=Inches(4.9),
                            font_size=14.5),
        caption="Cyan = new today  ·  DECISIONS sits right above RUN — "
                "it decides which version of the program actually "
                "happens")
    set_notes(slide, markers=["LIVE"], timing="8 min (Guide §19) — "
              "NON-NEGOTIABLE, HERO MOMENT",
              purpose="Build this top to bottom, one line at a time - "
                      "never reveal it finished. DECISIONS is the last "
                      "thing that happens before RUN.",
              board="Whiteboard - the one non-screen drawing before the "
                    "activity.",
              transition="Time to try this yourselves.")
    return slide


def slide_14_activity_intro(prs):
    slide = layout_activity(
        prs, "Try It Yourself", "Your Profile Card Decides",
        ["Add an if/else using is_student,", "and an if/elif/else chain "
         "using age", "with at least three branches."],
        sub="Change one value and re-run — confirm a different branch "
            "executes.")
    set_notes(slide, markers=["ACTIVITY"], timing="15 min (Guide §20) "
              "— NON-NEGOTIABLE, HERO ACTIVITY",
              purpose="Every student extends their own running Profile "
                      "Card with real branching for the first time. "
                      "Circulate constantly - check indentation first "
                      "for any error.",
              say="Stuck? What are two different messages that could "
                  "depend on whether you're a student? What are three "
                  "age categories you could sort yourself into?",
              board="Circulate and read screens rather than calling on "
                    "students verbally.",
              transition="Here's a template if you want a starting "
                          "point.")
    return slide


def slide_15_activity_template(prs):
    slide = layout_diagram(
        prs, "A Starting Point (Optional)", "Profile Card Decides — Template",
        draw_code_block(
            ["if is_student:", '    print("Time to study!")', "else:",
             '    print("Time to relax!")', "",
             "if age < 13:", '    print("Category: Child")',
             "elif age < 20:", '    print("Category: Teen")', "else:",
             '    print("Category: Adult")'], font_size=13.5))
    set_notes(slide, timing="(support for Guide §20)",
              purpose="A worked template for students who want a "
                      "starting point - not the only correct answer, "
                      "and not required.",
              board="None.",
              transition="Let a few students share one conditional and "
                          "both of its outcomes with the room.")
    return slide


def slide_16_activity_question(prs):
    slide = layout_big_question(
        prs, "Reflect",
        "What changed when you edited\none variable and ran your\n"
        "program again?", size=32,
        note="A different branch executed — same code, new data.")
    set_notes(slide, markers=["ASK", "PAUSE"], timing="(closes Guide "
              "§20)",
              purpose="Re-confirm Section 11's lesson using the "
                      "student's own program as the example.",
              board="None.",
              transition="Let's check a couple of misconceptions "
                          "before we close.")
    return slide


def slide_17_recap(prs):
    slide = new_slide(prs, bg=LIGHT_BG)
    y = _content_header(slide, "Quick Recap", "Quick Recap")
    draw_word_grid(["CONDITIONAL", "INDENTATION", "IF", "ELIF", "ELSE",
                     "BRANCH"], cols=6, font_size=12.5)(
        slide, MARGIN, y + Inches(0.1), CONTENT_W, Inches(1.1))
    draw_pipeline(
        [("PROBLEM", None), ("LOGIC", None), ("ALGORITHM", None),
         ("PYTHON\nCODE", None), ("VARIABLES", None), ("EXPRESSIONS", None),
         ("DECISIONS", None), ("RUN", None), ("OUTPUT/\nRESULT", None)],
        height=Inches(1.15), font_size=8.7, gap=Inches(0.18))(
        slide, MARGIN, y + Inches(1.7), CONTENT_W, Inches(1.4))
    add_footer(slide)
    set_notes(slide, markers=["ASK"], timing="5 min (Guide §22) — "
              "NON-NEGOTIABLE",
              purpose="One visual summary of every keyword and the "
                      "full bridged chain.",
              question="When does an if block's indented line run? "
                        "Why does indentation matter? When does else "
                        "run? If more than one elif condition is true, "
                        "which one runs? Can a condition use and/or?",
              board="None.",
              transition="One question to leave you with.")
    return slide


def slide_18_bridge_forward(prs):
    slide = layout_big_question(
        prs, "A Question for Next Time",
        "Right now, every decision only runs\nonce. What if a program "
        "needed to\nmake the same decision — or repeat\nthe same "
        "action — over and over?", size=27)
    set_notes(slide, markers=["DO NOT ANSWER"], timing="(Guide §23)",
              purpose="Leave this open as a deliberate hook into Class "
                      "09 (Repetition: while and for Loops).",
              board="None.",
              transition="Class ends on the final takeaway.")
    return slide


def slide_19_final_takeaway(prs):
    slide = layout_section_divider(
        prs, "Final Takeaway",
        "“Every program before today ran\nthe same way, every single "
        "time. Today,\nfor the first time, your code actually\nthinks "
        "before it acts — it looks at its\nown data and chooses a "
        "path.”",
        title_size=25)
    set_notes(slide, timing="2 min",
              purpose="Close the class on this exact line, said slowly.",
              board="None.",
              transition="Class ends.")
    return slide


# ============================================================================
# 6. MAIN GENERATION FUNCTION
# ============================================================================

SLIDE_BUILDERS = [
    slide_01_title, slide_02_big_question, slide_03_what_is_a_conditional,
    slide_04_first_if, slide_05_indentation, slide_06_else, slide_07_elif,
    slide_08_order_question, slide_09_order_reveal,
    slide_10_comparisons_logic, slide_11_common_patterns,
    slide_12_class01_callback, slide_13_full_bridge,
    slide_14_activity_intro, slide_15_activity_template,
    slide_16_activity_question, slide_17_recap, slide_18_bridge_forward,
    slide_19_final_takeaway,
]


def build_presentation(output_path):
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    for i, builder in enumerate(SLIDE_BUILDERS, start=1):
        slide = builder(prs)
        if i != 1:
            add_page_number(slide, i)
    prs.save(output_path)
    return len(SLIDE_BUILDERS)


if __name__ == "__main__":
    here = os.path.dirname(os.path.abspath(__file__))
    out = os.path.join(here, "Class_08_Making_Decisions.pptx")
    n = build_presentation(out)
    print(f"Generated {n} slides -> {out}")
